"""PR121 canonical data-integrity regressions; synthetic projects only."""
from copy import deepcopy
from datetime import date
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from reportmaker.model import HeadingBlock, LocalizedText, Module, ParagraphBlock, Report, TableBlock
from reportmaker.project_io import ReportProject, load_project
from reportmaker.reference import manual_snapshot
from reportmaker.rich_text import ParagraphStyle, RichFormatting, TextRun
from reportmaker.v2.adapters.persistence import ProjectPersistence
from reportmaker.v2.commands import CommandGateway, CommandIntent
from reportmaker.v2.session import ReportSession, StaleIntentError
from reportmaker.v2.web_edits import field_intent, rich_format_intent, table_shape_intent


def sample():
    text = LocalizedText('en', 'ABCDE')
    text.set_translation('it', 'Traduzione')
    cell = LocalizedText('en', 'Cell')
    cell.set_translation('it', 'Cella')
    rich = RichFormatting([TextRun(0, 5, bold=True)])
    module = Module(LocalizedText('en', 'Module'), [
        HeadingBlock(deepcopy(text), level=2, rich=deepcopy(rich)),
        ParagraphBlock(deepcopy(text), rich=deepcopy(rich)),
        TableBlock([[deepcopy(cell), deepcopy(cell)], [deepcopy(cell), deepcopy(cell)]])])
    return ReportProject(Report(customer_id=None, customer_snapshot=manual_snapshot('Synthetic'),
        start_date=date(2026, 10, 30), end_date=date(2026, 10, 30),
        primary_language='en', output_languages=['en', 'it'], modules=[module],
        conclusion=Module(LocalizedText('en', 'Conclusions'))))


class WebCommandTests(unittest.TestCase):
    def setUp(self):
        self.session = ReportSession(sample())
        self.gateway = CommandGateway(self.session)
        self.module = self.session.report_snapshot().modules[0]

    def edit(self, method, *args):
        intent = field_intent(self.session, method, list(args), self.session.revision, self.session.generation)
        return self.gateway.execute(intent)

    def block(self, index):
        return self.session.report_snapshot().modules[0].blocks[index]

    def test_heading_level_only_revision_and_undo(self):
        heading = self.module.blocks[0]
        self.edit('update_heading', self.module.id, heading.id, heading.text.original, 3)
        self.assertEqual(self.session.revision, 1)
        self.assertEqual(self.block(0).level, 3)
        self.assertEqual(self.block(0).text, heading.text)
        self.gateway.undo(self.session.generation, self.session.revision)
        self.assertEqual(self.block(0), heading)
        self.gateway.redo(self.session.generation, self.session.revision)
        self.assertEqual(self.block(0).level, 3)

    def test_heading_text_and_level_atomic_invalid_and_valid(self):
        heading = self.module.blocks[0]
        for level in (0, 7, True, '2'):
            with self.assertRaises(ValueError):
                self.edit('update_heading', self.module.id, heading.id, 'ABCDEF', level)
            self.assertEqual(self.block(0), heading)
            self.assertEqual(self.session.revision, 0)
        self.edit('update_heading', self.module.id, heading.id, 'ABCDEF', 3)
        self.assertEqual(self.session.revision, 1)
        self.assertEqual(self.block(0).rich.runs, heading.rich.runs)
        self.assertIn('it', self.block(0).text.stale_languages)
        self.gateway.undo(self.session.generation, self.session.revision)
        self.assertEqual(self.block(0), heading)

    def test_paragraph_rebases_unicode_and_preserves_translation(self):
        paragraph = self.module.blocks[1]
        self.edit('update_paragraph', self.module.id, paragraph.id, '🙂ABCDE!')
        edited = self.block(1)
        self.assertEqual(edited.rich.runs, [TextRun(1, 6, bold=True)])
        self.assertEqual(edited.text.translations, paragraph.text.translations)
        self.assertIn('it', edited.text.stale_languages)

    def test_ambiguous_styled_boundary_rejected(self):
        project = sample()
        block = project.report.modules[0].blocks[1]
        block.text.original = 'A\nB'
        block.rich = RichFormatting(paragraphs=[ParagraphStyle(2, align='center')])
        session = ReportSession(project)
        with self.assertRaisesRegex(ValueError, 'boundary'):
            field_intent(session, 'update_paragraph', [project.report.modules[0].id, block.id, 'AB'], 0, session.generation)
        self.assertEqual(session.revision, 0)

    def test_table_all_shape_changes_preserve_cells_and_undo(self):
        table = self.module.blocks[2]
        for op in ('add_row', 'add_col', 'del_row', 'del_col'):
            self.gateway.execute(table_shape_intent(self.session, self.module.id, table.id, op))
            self.assertEqual(self.block(2).rows[0][0], table.rows[0][0])
            self.gateway.undo(self.session.generation, self.session.revision)
            self.assertEqual(self.block(2), table)
            self.gateway.redo(self.session.generation, self.session.revision)
            self.gateway.undo(self.session.generation, self.session.revision)
        with self.assertRaises(ValueError):
            table_shape_intent(self.session, self.module.id, table.id, 'surprise')

    def test_table_edit_stales_only_changed_cell_and_roundtrip(self):
        table = self.module.blocks[2]
        self.edit('update_table', self.module.id, table.id, json.dumps([['New', 'Cell'], ['Cell', 'Cell']]))
        edited = self.block(2)
        self.assertEqual(edited.rows[0][0].translations, table.rows[0][0].translations)
        self.assertEqual(edited.rows[0][0].stale_languages, {'it'})
        self.assertEqual(edited.rows[0][1], table.rows[0][1])
        self.gateway.execute(table_shape_intent(self.session, self.module.id, table.id, 'add_row'))
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'roundtrip.rmproj'
            ProjectPersistence().save(self.session, path)
            reopened = load_project(path)
            self.assertEqual(reopened.report, self.session.report_snapshot())
        self.gateway.undo(self.session.generation, self.session.revision)
        self.assertEqual(self.block(2), edited)

    def test_bad_table_values_do_not_coerce_or_reshape(self):
        table = self.module.blocks[2]
        for raw in ('{}', '[["only"]]', '[[1,"Cell"],["Cell","Cell"]]'):
            with self.assertRaises(ValueError):
                self.edit('update_table', self.module.id, table.id, raw)
            self.assertEqual(self.block(2), table)

    def test_atomic_field_batch_with_metadata_and_invalid_tail(self):
        session = self.session
        heading = self.module.blocks[0]
        def build(level):
            edits = [field_intent(session, 'update_report_header', [json.dumps({'ticket': 'changed'})], 0, session.generation),
                     field_intent(session, 'update_heading', [self.module.id, heading.id, 'ABCDEF', level], 0, session.generation)]
            return CommandIntent('set_fields_batch', None, {'edits': edits}, 0, session.generation)
        original = session.report_snapshot()
        with self.assertRaises(ValueError):
            self.gateway.execute(build(99))
        self.assertEqual(session.report_snapshot(), original)
        self.gateway.execute(build(3))
        self.assertEqual(session.revision, 1)
        self.gateway.undo(session.generation, session.revision)
        self.assertEqual(session.report_snapshot(), original)

    def test_stale_envelope_never_retargets_previous_project(self):
        heading = self.module.blocks[0]
        intent = field_intent(self.session, 'update_heading', [self.module.id, heading.id, 'ABCDEF', 3], 0, self.session.generation)
        self.session.replace_project(sample(), discard_unsaved=True)
        with self.assertRaises(StaleIntentError):
            self.gateway.execute(intent)

    def test_manual_snapshot_details_survive_header_edit(self):
        project = sample()
        project.report.customer_snapshot = manual_snapshot('Synthetic', customer_address='Retained address')
        self.session = ReportSession(project)
        self.gateway = CommandGateway(self.session)
        self.edit('update_report_header', json.dumps({'title': 'Changed title'}))
        self.assertEqual(self.session.report_snapshot().customer_snapshot, project.report.customer_snapshot)

    def test_table_shape_roundtrip_multilingual_word_and_pdf(self):
        from reportmaker.v2.localized_export import localized_fields, localized_project, TranslationRequiredError
        from reportmaker.export import export_project
        from docx import Document
        from pypdf import PdfReader
        project = sample()
        for _, text in localized_fields(project.report):
            text.set_translation('it', 'Cella' if text.original == 'Cell' else 'Testo italiano')
        self.session = ReportSession(project)
        self.gateway = CommandGateway(self.session)
        module = project.report.modules[0]
        table = module.blocks[2]
        self.gateway.execute(table_shape_intent(self.session, module.id, table.id, 'add_row'))
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'translated.rmproj'
            ProjectPersistence().save(self.session, path)
            reopened = load_project(path)
            italian = localized_project(reopened, 'it')
            word = export_project(italian, Path(directory)/'italian.docx', format='docx')
            pdf = export_project(italian, Path(directory)/'italian.pdf', format='pdf')
            self.assertTrue(any(cell.text == 'Cella' for t in Document(word).tables for row in t.rows for cell in row.cells))
            self.assertIn('Cella', ''.join(p.extract_text() for p in PdfReader(pdf).pages))
        self.edit('update_table', module.id, table.id, json.dumps([['Changed', 'Cell'], ['Cell', 'Cell'], ['', '']]))
        with self.assertRaises(TranslationRequiredError):
            localized_project(self.session.project_snapshot(), 'it')

    def test_word_ribbon_rich_format_intent_operations(self):
        para = self.module.blocks[1]
        for spec in (
            {'op': 'bold', 'start': 0, 'end': 3},
            {'op': 'italic', 'start': 0, 'end': 3},
            {'op': 'underline', 'start': 0, 'end': 3},
            {'op': 'strike', 'start': 0, 'end': 3},
            {'op': 'baseline', 'value': 'sub', 'start': 0, 'end': 2},
            {'op': 'family', 'value': 'Calibri Light', 'start': 0, 'end': 6},
            {'op': 'size', 'value': 18, 'start': 0, 'end': 6},
            {'op': 'step_size', 'value': 1, 'start': 0, 'end': 6},
            {'op': 'color', 'value': '#1d4ed8', 'start': 0, 'end': 6},
            {'op': 'highlight', 'value': '#ffff00', 'start': 0, 'end': 6},
            {'op': 'align', 'value': 'center', 'start': 0, 'end': 0},
            {'op': 'list_kind', 'value': 'bullet', 'start': 0, 'end': 0},
            {'op': 'indent', 'value': 1, 'start': 0, 'end': 0},
            {'op': 'line_height', 'value': 1.5, 'start': 0, 'end': 0},
            {'op': 'shading', 'value': '#dbeafe', 'start': 0, 'end': 0},
            {'op': 'border', 'value': 'toggle', 'start': 0, 'end': 0},
            {'op': 'change_case', 'value': 'lower', 'start': 0, 'end': 5},
        ):
            intent = rich_format_intent(
                self.session, self.module.id, para.id,
                json.dumps(spec), self.session.revision, self.session.generation,
            )
            self.gateway.execute(intent)
        updated = self.session.report_snapshot().modules[0].blocks[1]
        self.assertEqual(updated.text.original, 'abcde')
        self.assertFalse(updated.rich.runs[0].bold)
        self.assertTrue(updated.rich.runs[0].italic)
        self.assertTrue(updated.rich.runs[0].underline)
        self.assertTrue(updated.rich.runs[0].strike)
        self.assertEqual(updated.rich.runs[0].family, 'Calibri Light')
        self.assertEqual(updated.rich.runs[0].size, 20.0)
        self.assertEqual(updated.rich.paragraphs[0].align, 'center')
        self.assertEqual(updated.rich.paragraphs[0].list_kind, 'bullet')
        self.assertEqual(updated.rich.paragraphs[0].indent, 1)
        self.assertEqual(updated.rich.paragraphs[0].line_height, 1.5)
        self.assertEqual(updated.rich.paragraphs[0].shading, '#dbeafe')
        self.assertTrue(updated.rich.paragraphs[0].border)


if __name__ == '__main__':
    unittest.main()
