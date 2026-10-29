"""M6.7 Checkpoint C: synthetic, shared-session, user-path integration gates.

Qt runs offscreen on both Windows CI and local Linux. No WebEngine activation,
customer content, production images, network access or installed Office needed.
"""
from __future__ import annotations

from dataclasses import replace
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile
import json
import os
import unittest

from docx import Document
from PIL import Image
from pypdf import PdfReader

from reportmaker.engine import module_outline
from reportmaker.model import LocalizedText, PhotoMark, PhotoVisual
from reportmaker.project_io import load_project
from reportmaker.v2.commands import CommandIntent
from reportmaker.v2.localized_export import localized_fields
from reportmaker.v2.metadata_locale import MetadataLocale, visible_metadata_keys
from reportmaker.v2.projections import NodeAddress
from test_v2_m6 import synthetic_project


def execute(win, kind, args, address=None):
    return win.gateway.execute(CommandIntent(
        kind, address, args, win.session.revision, win.session.generation))


def approve_all(win, lang="de"):
    """Explicitly approve every selected export field via the canonical gateway."""
    report = win.session.report_snapshot()
    for key, value in localized_fields(report):
        if value.original.strip() and value.original_language.casefold() != lang.casefold():
            execute(win, "set_translation", {
                "key": key, "language": lang, "text": "DE · " + value.original})
    report = win.session.report_snapshot()
    locale = MetadataLocale(report)
    for key in visible_metadata_keys(report):
        source, source_language = locale.source(key)
        if source.strip() and source_language.casefold() != lang.casefold():
            approval = locale.approve(key, lang, "DE · " + source)
            execute(win, "set_metadata_translation", {
                "key": key, "language": lang,
                "text": approval.text,
                "expected_source_sha256": approval.source_sha256})


@unittest.skipUnless(os.environ.get("QT_QPA_PLATFORM") == "offscreen",
                     "Dedicated offline Qt Checkpoint C gate")
class CheckpointCQtTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        from reportmaker.v2.app import V2MainWindow
        self.win = V2MainWindow(
            project=synthetic_project(), load_website=False,
            ephemeral_browser=True)
        self.win.show()
        self.pump()

    def tearDown(self):
        self.win._confirm_transition = lambda: False
        self.win.close()
        self.pump()

    def pump(self):
        for _ in range(6):
            self.app.processEvents()

    def test_full_checkpoint_c_roundtrip_uses_one_report_and_real_export_paths(self):
        win = self.win
        report = win.session.report_snapshot()
        original_photo = win.session.assets["synthetic_photo"]
        text_module = report.modules[0].blocks[1].module
        text_block = text_module.blocks[0]
        table = report.modules[0].blocks[3]
        text_address = NodeAddress(report.id, text_module.id, text_block.id, "text")
        table_address = NodeAddress(report.id, report.modules[0].id, table.id, "table")
        photo_address = NodeAddress(report.id, report.modules[0].id,
                                    report.modules[0].blocks[2].id, "photo")

        with TemporaryDirectory() as directory:
            folder = Path(directory)
            picture = Image.new("RGB", (128, 92), (12, 90, 140))
            sample_path = folder / "extra-synthetic-photo.png"
            picture.save(sample_path)
            win.tabs.setCurrentWidget(win.gallery_tab)
            self.assertTrue(win.gallery_tab.import_paths([str(sample_path)]))
            self.assertEqual(len(win.session.media_library), 1)
            self.assertEqual(win.session.assets["synthetic_photo"], original_photo)

            # Report Information remains canonical; detached manual snapshots
            # never require a customer database/master record.
            updated_snapshot = replace(
                report.customer_snapshot, site_name="Fictional Test Site",
                machine_name="Laboratory Pump", machine_serial="TEST-0007")
            execute(win, "set_report_metadata", {
                "title": "Synthetic M6.7 service report",
                "created_date": report.created_date,
                "start_date": report.start_date, "end_date": report.end_date,
                "ticket": "SYN-CHECKPOINT-C",
                "technicians": ["Fictional Technician"],
                "intervention_types": ["supervision", "ordinary_maintenance"],
                "other_intervention_detail": "",
                "customer_snapshot": updated_snapshot,
                "customer_id": None,
                "primary_language": "en", "output_languages": ["en", "de"]})
            self.assertIsNone(win.session.report_snapshot().customer_id)

            win.tabs.setCurrentWidget(win.editor_tab)
            self.pump()
            win.editor_tab._editors[text_address].setPlainText(
                "Synthetic bearing inspection accepted.")
            win.editor_tab._tables[table_address].item(0, 0).setText(
                "Inspection value")
            self.pump()
            self.assertTrue(win.branch.drafts.pending)
            self.assertEqual(win.branch.drafts.draft_owner(text_address),
                             "page-editor")
            self.assertEqual(win.session.report_snapshot().modules[0].blocks[1]
                             .module.blocks[0].text.original, "Torque checked")
            with patch("reportmaker.v2.app.QMessageBox.warning"):
                self.assertFalse(win.save_project(folder / "pending.rmproj"))
            self.assertFalse(win.preview_tab.export_to(folder / "pending.pdf", "pdf"))
            self.assertFalse(win.preview_tab.export_bundle_to(folder / "pending.zip"))
            self.assertFalse((folder / "pending.rmproj").exists())
            self.assertFalse((folder / "pending.pdf").exists())
            self.assertFalse((folder / "pending.zip").exists())

            # One atomic Apply owns both page fields; Undo and Redo leave
            # the same stable IDs, translations and report outline intact.
            before = win.session.report_snapshot()
            rev = win.session.revision
            self.assertTrue(win.editor_tab.apply_changes())
            self.assertEqual(win.session.revision, rev + 1)
            approved = win.session.report_snapshot()
            self.assertEqual(approved.modules[0].blocks[1].module.blocks[0]
                             .text.original, "Synthetic bearing inspection accepted.")
            self.assertEqual(approved.modules[0].blocks[3].rows[0][0].original,
                             "Inspection value")
            self.assertTrue(win.branch.undo())
            self.assertEqual(win.session.report_snapshot(), before)
            self.assertTrue(win.branch.redo())
            self.assertEqual(win.session.report_snapshot(), approved)

            execute(win, "set_photo_visual", {
                "visual": PhotoVisual(rotation=90, marks=[
                    PhotoMark("arrow", .1, .2, .8, .7),
                    PhotoMark("text", .2, .3, .6, .6,
                              text=LocalizedText("en", "Synthetic marker"))])},
                photo_address)

            # Explicit language approvals are commands, not fixture-side
            # mutations or a parallel copy of report content.
            approve_all(win)
            finalized = win.session.report_snapshot()
            identity = [(entry.module_id, entry.number, entry.depth)
                        for entry in module_outline(finalized)]
            self.assertEqual([action.data() for action in win.build_contents_menu().actions()],
                             [entry.module_id for entry in module_outline(finalized)])
            self.assertIs(win.editor_tab.session, win.session)
            self.assertIs(win.gallery_tab.session, win.session)
            self.assertIs(win.preview_tab.session, win.session)
            self.assertIsNone(win.chat_tab.browser)
            self.assertFalse(win.branch.drafts.pending)

            project_path = folder / "checkpoint-c.rmproj"
            self.assertTrue(win.save_project(project_path))
            self.assertFalse(win.session.dirty)
            saved = load_project(project_path)
            self.assertEqual(saved.report, finalized)
            self.assertEqual(saved.assets, dict(win.session.assets))
            self.assertEqual(saved.media_library, dict(win.session.media_library))

            # Open through the actual main-window lifecycle and ensure every
            # workspace is still bound to the one newly installed generation.
            previous_generation = win.session.generation
            self.assertTrue(win.open_project(project_path))
            self.pump()
            self.assertNotEqual(win.session.generation, previous_generation)
            self.assertEqual(win.session.report_snapshot(), finalized)
            self.assertEqual(
                [(entry.module_id, entry.number, entry.depth)
                 for entry in module_outline(win.session.report_snapshot())], identity)
            self.assertEqual(win.session.assets["synthetic_photo"], original_photo)

            win.tabs.setCurrentWidget(win.preview_tab)
            self.pump()
            self.assertTrue(win.preview_tab.export_to(folder / "original.docx", "docx"))
            self.assertTrue(win.preview_tab.export_to(folder / "original.pdf", "pdf"))
            win.preview_tab.language.setCurrentIndex(
                win.preview_tab.language.findData("de"))
            self.pump()
            self.assertTrue(win.preview_tab.show_pdf_proof())
            self.assertTrue(win.preview_tab.export_to(folder / "translated.docx", "docx"))
            self.assertTrue(win.preview_tab.export_bundle_to(folder / "all-languages.zip"))
            self.assertEqual(win.session.report_snapshot(), finalized)
            doc = Document(folder / "original.docx")
            translated = Document(folder / "translated.docx")
            all_source = "\n".join(p.text for p in doc.paragraphs)
            all_de = "\n".join(p.text for p in translated.paragraphs)
            self.assertIn("Synthetic bearing inspection accepted.", all_source)
            self.assertIn("DE · Synthetic bearing inspection accepted.", all_de)
            self.assertEqual(doc.tables[0].cell(0, 0).text, "Created")
            self.assertIn("SYN-CHECKPOINT-C", "\n".join(
                cell.text for tbl in doc.tables for row in tbl.rows for cell in row.cells))
            self.assertGreaterEqual(len(doc.inline_shapes), 1)
            self.assertGreaterEqual(len(doc.tables), 2)
            self.assertGreaterEqual(len(PdfReader(folder / "original.pdf").pages), 1)
            with ZipFile(folder / "all-languages.zip") as archive:
                self.assertEqual(set(archive.namelist()), {
                    "manifest.json", "report_en.docx", "report_de.docx",
                    "report_en.pdf", "report_de.pdf"})
                self.assertEqual(json.loads(archive.read("manifest.json"))["languages"],
                                 ["en", "de"])
                self.assertEqual(
                    archive.read("report_en.docx")[:2], b"PK")
                self.assertTrue(archive.read("report_de.pdf").startswith(b"%PDF"))
            with ZipFile(project_path) as packed:
                self.assertFalse(any(
                    term in member.lower() for member in packed.namelist()
                    for term in ("browser", "cookie", "chatgpt", "cache")))

            # A late source edit invalidates previously approved translations.
            # Both export paths must fail without overwriting approved files,
            # even though the live report remains editable and undoable.
            archive_before = (folder / "all-languages.zip").read_bytes()
            translated_before = (folder / "translated.docx").read_bytes()
            execute(win, "set_paragraph", {
                "text": "Unapproved new source words"}, text_address)
            self.assertFalse(win.preview_tab.export_bundle_to(
                folder / "all-languages.zip"))
            self.assertFalse(win.preview_tab.export_to(
                folder / "translated.docx", "docx"))
            self.assertEqual((folder / "all-languages.zip").read_bytes(),
                             archive_before)
            self.assertEqual((folder / "translated.docx").read_bytes(),
                             translated_before)
            self.assertTrue(win.branch.undo())
            self.assertEqual(win.session.report_snapshot(), finalized)

    def test_failed_reload_and_stale_language_export_preserve_saved_bytes(self):
        win = self.win
        with TemporaryDirectory() as directory:
            base = Path(directory)
            path = base / "good.rmproj"
            self.assertTrue(win.save_project(path))
            preserved_report = win.session.report_snapshot()
            preserved_generation = win.session.generation
            self.assertFalse(win.windowTitle().endswith(" *"))
            saved_report = win.session.report_snapshot()
            child = saved_report.modules[0].blocks[1].module
            text_address = NodeAddress(saved_report.id, child.id,
                                        child.blocks[0].id, "text")
            win.tabs.setCurrentWidget(win.editor_tab)
            self.pump()
            win.editor_tab._editors[text_address].setPlainText("Pending only")
            self.pump()
            self.assertTrue(win.branch.drafts.pending)
            self.assertTrue(win.windowTitle().endswith(" *"))
            self.assertTrue(win.editor_tab.revert_changes())
            self.pump()
            self.assertFalse(win.windowTitle().endswith(" *"))
            win.branch.drafts.update(text_address, "Unapplied branch wording")
            self.assertTrue(win.windowTitle().endswith(" *"))
            win.branch.drafts.discard(text_address)
            self.assertFalse(win.windowTitle().endswith(" *"))
            corrupt = base / "corrupt.rmproj"
            corrupt.write_bytes(b"Invalid synthetic project, no user data")
            with patch("reportmaker.v2.app.QMessageBox.critical"):
                self.assertFalse(win.open_project(corrupt))
            self.assertEqual(win.session.generation, preserved_generation)
            self.assertEqual(win.session.report_snapshot(), preserved_report)
            self.assertEqual(win.session.path, path)
            archived = base / "preexisting.zip"
            archived.write_bytes(b"trusted original")
            self.assertFalse(win.preview_tab.export_bundle_to(archived))
            self.assertEqual(archived.read_bytes(), b"trusted original")
            self.assertEqual(win.session.report_snapshot(), preserved_report)
            self.assertEqual(sorted(x.name for x in base.iterdir()),
                             ["corrupt.rmproj", "good.rmproj", "preexisting.zip"])


if __name__ == "__main__":
    unittest.main()
