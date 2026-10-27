"""Synthetic-only M4 export regressions. No production reports or web access."""
from __future__ import annotations

from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile

from docx import Document
from PIL import Image
from pypdf import PdfReader

from reportmaker.engine import module_outline
from reportmaker.export import ExportError, export_project, prepare_export
from reportmaker.project_io import MAX_ASSET
from reportmaker.model import (
    HeadingBlock, LocalizedText, Module, NestedModuleBlock, PageBreakBlock,
    ParagraphBlock, PhotoBlock, Report, TableBlock,
)
from reportmaker.project_io import ReportProject, save_project
from reportmaker.reference import manual_snapshot


def loc(value):
    return LocalizedText("it", value)


def synthetic_project():
    picture = Image.new("RGB", (20, 12), (11, 88, 140))
    raw = BytesIO()
    picture.save(raw, "PNG")
    child = Module(loc("Controllo interno"), [
        ParagraphBlock(loc("Pressione stabile\nSeconda riga")),
        PhotoBlock("synthetic", loc("Foto dimostrativa")),
        TableBlock([[loc("Parametro"), loc("Valore")],
                    [loc("Velocità"), loc("45")]]),
        PageBreakBlock(),
        HeadingBlock(loc("Dettagli"), 2),
    ])
    first = Module(loc("Verifica linea"), [
        ParagraphBlock(loc("Verifica avviata")),
        NestedModuleBlock(child),
        ParagraphBlock(loc("Nota dopo il sottomodulo")),
    ])
    last = Module(loc("Prova"), [ParagraphBlock(loc("Controllo completato"))])
    report = Report(
        customer_id=None, customer_snapshot=manual_snapshot(
            "Azienda Fittizia", site_name="Stabilimento di test",
            machine_name="Macchina campione", machine_serial="SYN-0042"),
        created_date=date(2026, 10, 26),
        start_date=date(2026, 10, 24), end_date=date(2026, 10, 26),
        primary_language="it", output_languages=["it", "de-AT"],
        title=loc("Rapporto tecnico"), technicians=["Tecnico Uno", "Tecnico Due"],
        intervention_types=["supervision", "extraordinary_maintenance", "other"],
        other_intervention_detail="Prova funzionale sintetica",
        modules=[first, last],
        conclusion=Module(loc("Conclusioni"), [ParagraphBlock(loc("Esito positivo"))]),
    )
    return ReportProject(report, {"synthetic": raw.getvalue()})


class M4ExportTests(unittest.TestCase):
    def test_single_outline_and_frozen_metadata(self):
        project = synthetic_project()
        plan = prepare_export(project)
        self.assertEqual(plan.outline, module_outline(project.report))
        self.assertEqual(
            [(item.kind, item.text) for item in plan.items if item.kind == "module"],
            [("module", "1 Verifica linea"), ("module", "1.1 Controllo interno"),
             ("module", "2 Prova"), ("module", "Conclusioni")])
        self.assertEqual(plan.items[-1].text, "Esito positivo")
        self.assertIn(("Machine serial", "SYN-0042"), plan.info)
        self.assertIn(("Other intervention details", "Prova funzionale sintetica"), plan.info)
        self.assertIn(("Intervention types",
                       "Supervision, Extraordinary maintenance, Other"), plan.info)
        # Arbitrary language identifiers remain data, not UI enums.
        self.assertEqual(project.report.output_languages, ["it", "de-AT"])

    def test_editable_word_elements_toc_refs_and_mixed_order(self):
        project = synthetic_project()
        before = deepcopy(project)
        with TemporaryDirectory() as directory:
            path = Path(directory) / "test.docx"
            self.assertEqual(export_project(project, path, format="docx"), path)
            document = Document(path)
            paragraphs = [p.text for p in document.paragraphs]
            self.assertIn("Rapporto tecnico", paragraphs)
            self.assertIn("1.1 Controllo interno", paragraphs)
            self.assertIn("Nota dopo il sottomodulo", paragraphs)
            self.assertGreater(paragraphs.index("Nota dopo il sottomodulo"),
                               paragraphs.index("1.1 Controllo interno"))
            self.assertGreater(paragraphs.index("Conclusioni"),
                               paragraphs.index("2 Prova"))
            self.assertIn("Esito positivo", paragraphs)
            self.assertGreaterEqual(len(document.tables), 2)
            self.assertEqual(document.tables[-1].cell(1, 0).text, "Velocità")
            self.assertEqual(document.tables[-1].cell(1, 1).text, "45")
            self.assertEqual(len(document.inline_shapes), 1)
            with ZipFile(path) as archive:
                xml = archive.read("word/document.xml").decode()
                settings = archive.read("word/settings.xml").decode()
            self.assertEqual(xml.count("PAGEREF"), len(module_outline(project.report)))
            self.assertIn("bookmarkStart", xml)
            self.assertIn("updateFields", settings)
            self.assertIn("Prova funzionale sintetica", xml)
            self.assertEqual(project, before)

    def test_pdf_pagination_contents_and_photo(self):
        project = synthetic_project()
        with TemporaryDirectory() as directory:
            path = Path(directory) / "test.pdf"
            export_project(project, path, format="pdf")
            reader = PdfReader(path)
            self.assertGreaterEqual(len(reader.pages), 2)
            pages = [page.extract_text() for page in reader.pages]
            text = "\n".join(pages)
            self.assertIn("Rapporto tecnico", text)
            self.assertIn("Controllo interno", text)
            self.assertIn("Nota dopo il sottomodulo", text)
            self.assertIn("Conclusioni", text)
            self.assertIn("Esito positivo", text)
            self.assertIn("Prova funzionale sintetica", text)
            self.assertIn("Page 1", pages[0])
            self.assertIn("Summary", pages[0])
            self.assertGreaterEqual(len(reader.outline), 3)

    def test_failed_export_cannot_replace_existing_file(self):
        project = synthetic_project()
        with TemporaryDirectory() as directory:
            for kind in ("docx", "pdf"):
                output = Path(directory) / ("unchanged." + kind)
                output.write_bytes(b"original user document")
                bad = deepcopy(project)
                bad.assets["synthetic"] = b"not an image"
                with self.assertRaises(ExportError):
                    export_project(bad, output, format=kind)
                self.assertEqual(output.read_bytes(), b"original user document")
                bad = deepcopy(project)
                del bad.assets["synthetic"]
                with self.assertRaises(ExportError):
                    export_project(bad, output, format=kind)
                self.assertEqual(output.read_bytes(), b"original user document")
                with self.assertRaises(ExportError):
                    export_project(project, output, format="unknown")
                self.assertEqual(output.read_bytes(), b"original user document")
            self.assertFalse(list(Path(directory).glob(".reportmaker-export-*")))

    def test_pdf_toc_escapes_literal_module_titles_and_unicode_metadata(self):
        project = synthetic_project()
        project.report.modules[0].title.update_original("Valve <on> & motor")
        project.report.customer_snapshot = manual_snapshot("Клиент Тест")
        with TemporaryDirectory() as directory:
            pdf = Path(directory) / "literal-titles.pdf"
            export_project(project, pdf, format="pdf")
            pages = [page.extract_text() for page in PdfReader(pdf).pages]
            # The first page is the generated Summary, not the body heading.
            self.assertIn("Valve <on> & motor", pages[0])
            self.assertIn("Valve <on> & motor", "\n".join(pages[1:]))
            # The selected Segoe UI/DejaVu font must also cover metadata.
            from os import environ
            installed_fonts = (
                (Path(environ.get("WINDIR", "C:/Windows")) / "Fonts" / "segoeui.ttf").is_file()
                or Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf").is_file()
            )
            if installed_fonts:
                self.assertIn("Клиент Тест", pages[0])

    def test_oversized_photo_rejected_before_decoder_and_no_clobber(self):
        project = synthetic_project()
        project.assets["synthetic"] = b"x" * (MAX_ASSET + 1)
        with TemporaryDirectory() as directory:
            target = Path(directory) / "existing.docx"
            target.write_bytes(b"saved output")
            with self.assertRaisesRegex(ExportError, "asset size limit"):
                export_project(project, target, format="docx")
            self.assertEqual(target.read_bytes(), b"saved output")

    def test_pdf_oversized_table_cell_produces_user_error_without_clobber(self):
        project = synthetic_project()
        project.report.modules[0].blocks.append(
            TableBlock([[loc("Too tall\n" * 350)]]))
        with TemporaryDirectory() as directory:
            target = Path(directory) / "existing.pdf"
            target.write_bytes(b"saved output")
            with self.assertRaisesRegex(ExportError, "does not fit"):
                export_project(project, target, format="pdf")
            self.assertEqual(target.read_bytes(), b"saved output")
            self.assertFalse(list(Path(directory).glob(".reportmaker-export-*")))

    def test_source_rmproj_does_not_mutate_and_cannot_be_output(self):
        project = synthetic_project()
        with TemporaryDirectory() as directory:
            source = Path(directory) / "saved.rmproj"
            save_project(project, source)
            original = source.read_bytes()
            with self.assertRaises(ExportError):
                export_project(project, source, format="docx")
            self.assertEqual(source.read_bytes(), original)
            export_project(project, Path(directory) / "copy.docx", format="docx")
            export_project(project, Path(directory) / "copy.pdf", format="pdf")
            self.assertEqual(source.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
