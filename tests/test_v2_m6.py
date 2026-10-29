"""M6 workspace integration regressions: synthetic records and images only."""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import os
import unittest

from PIL import Image
from reportmaker.model import PhotoMark, PhotoVisual
from reportmaker.export import ExportError, export_project, prepare_export
from reportmaker.v2.commands import CommandGateway, CommandIntent
from reportmaker.v2.projections import NodeAddress
from reportmaker.v2.session import PendingDraftError, ReportSession
from test_v2_structure import fixture


def synthetic_project():
    project = fixture()
    image = Image.new("RGB", (84, 56), (220, 225, 230))
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    project.assets["synthetic_photo"] = buffer.getvalue()
    return project


class IntegratedMediaTests(unittest.TestCase):
    def test_gallery_only_import_is_atomic_undoable_and_persistent(self):
        project = synthetic_project()
        session = ReportSession(project)
        gateway = CommandGateway(session)
        baseline = session.report_snapshot()
        sample = project.assets["synthetic_photo"]
        command = CommandIntent("import_library_photos", None,
                                {"photos": (("Synthetic unplaced", sample),)},
                                session.revision, session.generation)
        change = gateway.execute(command)
        self.assertTrue(change.assets_changed)
        self.assertFalse(change.structure_changed)
        self.assertEqual(session.report_snapshot(), baseline)
        self.assertEqual(len(session.media_library), 1)
        key = next(iter(session.media_library))
        self.assertEqual(session.assets[key], sample)
        self.assertTrue(gateway.undo(session.generation, session.revision))
        self.assertEqual(session.report_snapshot(), baseline)
        self.assertNotIn(key, session.assets)
        self.assertTrue(gateway.redo(session.generation, session.revision))
        self.assertEqual(session.assets[key], sample)
        with TemporaryDirectory() as directory:
            from reportmaker.v2.adapters.persistence import ProjectPersistence
            from reportmaker.project_io import load_project
            target = Path(directory) / "media.rmproj"
            ProjectPersistence().save(session, target)
            reopened = load_project(target)
            self.assertEqual(reopened.media_library[key], "Synthetic unplaced")
            self.assertEqual(reopened.assets[key], sample)

    def test_metadata_is_canonical_atomic_multitype_and_open_language(self):
        from datetime import date
        from dataclasses import replace
        project = synthetic_project()
        session = ReportSession(project)
        gateway = CommandGateway(session)
        baseline = session.report_snapshot()
        args = {
            "title": "Synthetic intervention summary",
            "created_date": date(2026, 10, 29),
            "start_date": date(2026, 10, 27),
            "end_date": date(2026, 10, 29),
            "ticket": "FAKE-123",
            "technicians": ["Synthetic Technician"],
            "intervention_types": ["supervision", "ordinary_maintenance", "other"],
            "other_intervention_detail": "Fabricated",
            "customer_snapshot": replace(
                baseline.customer_snapshot, site_name="Synthetic site",
                machine_name="Synthetic equipment"),
            "primary_language": "en",
            "output_languages": ["en", "de", "fr-CA"],
        }
        action = CommandIntent("set_report_metadata", None, args,
                               session.revision, session.generation)
        change = gateway.execute(action)
        self.assertTrue(change.outline_invalidated)
        report = session.report_snapshot()
        self.assertEqual(report.id, baseline.id)
        self.assertEqual(report.modules, baseline.modules)
        self.assertEqual(report.conclusion, baseline.conclusion)
        self.assertEqual(report.ticket, "FAKE-123")
        self.assertEqual(report.intervention_types, args["intervention_types"])
        self.assertEqual(report.output_languages, ["en", "de", "fr-CA"])
        with TemporaryDirectory() as directory:
            from reportmaker.v2.adapters.persistence import ProjectPersistence
            from reportmaker.project_io import load_project
            target = Path(directory) / "metadata.rmproj"
            ProjectPersistence().save(session, target)
            self.assertEqual(load_project(target).report, report)
        self.assertTrue(gateway.undo(session.generation, session.revision))
        self.assertEqual(session.report_snapshot(), baseline)
        self.assertTrue(gateway.redo(session.generation, session.revision))
        self.assertEqual(session.report_snapshot(), report)

    def test_invalid_report_information_preserves_revision_and_undo_history(self):
        from datetime import date
        project = synthetic_project()
        session = ReportSession(project)
        gateway = CommandGateway(session)
        original = session.report_snapshot()
        info = {
            "title": original.title.original,
            "created_date": original.created_date,
            "start_date": date(2026, 10, 29),
            "end_date": date(2026, 10, 28),
            "ticket": original.ticket,
            "technicians": original.technicians,
            "intervention_types": original.intervention_types,
            "other_intervention_detail": original.other_intervention_detail,
            "customer_snapshot": original.customer_snapshot,
            "primary_language": original.primary_language,
            "output_languages": original.output_languages,
        }
        with self.assertRaises(ValueError):
            gateway.execute(CommandIntent(
                "set_report_metadata", None, info, session.revision, session.generation))
        self.assertEqual(session.revision, 0)
        self.assertEqual(session.report_snapshot(), original)
        self.assertFalse(session._editor.can_undo)

    def test_pure_export_still_fails_closed_without_explicit_visual_renderer(self):
        project = synthetic_project()
        photo = project.report.modules[0].blocks[2]
        photo.visual = PhotoVisual(
            rotation=90,
            marks=[PhotoMark("arrow", .1, .1, .8, .8)],
        )
        with self.assertRaises(ExportError):
            prepare_export(project)
        from reportmaker.v2.visual_export import annotated_png
        plan = prepare_export(project, visual_renderer=annotated_png)
        photo_item = next(item for item in plan.items if item.kind == "photo")
        self.assertTrue(photo_item.image[:4] == bytes((137, 80, 78, 71)))
        self.assertNotEqual(photo_item.image, project.assets["synthetic_photo"])
        before = deepcopy(project.report)
        with TemporaryDirectory() as directory:
            for format in ("docx", "pdf"):
                target = Path(directory) / ("report." + format)
                self.assertEqual(
                    export_project(project, target, format=format,
                                   visual_renderer=annotated_png), target)
                self.assertGreater(target.stat().st_size, 500)
        self.assertEqual(project.report, before)


@unittest.skipUnless(os.environ.get("QT_QPA_PLATFORM") == "offscreen",
                     "Windows hosted offscreen Qt gate")
class IntegratedWorkspaceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        from reportmaker.v2.app import V2MainWindow
        self.project = synthetic_project()
        self.before = deepcopy(self.project.report)
        self.win = V2MainWindow(project=self.project, load_website=False,
                                ephemeral_browser=True)
        self.win.show()
        self.pump()

    def tearDown(self):
        self.win._confirm_transition = lambda: False
        self.win.close()
        self.pump()

    def pump(self):
        for _ in range(5):
            self.app.processEvents()

    def test_tabs_are_separate_views_of_one_session_and_webengine_is_lazy(self):
        tabs = self.win.tabs
        self.assertEqual(
            [tabs.tabText(i) for i in range(tabs.count())],
            ["Modules", "Editor", "Photo Gallery", "ChatGPT", "Preview"])
        self.assertIsNone(self.win.chat_tab.browser)
        tabs.setCurrentWidget(self.win.preview_tab)
        self.win.preview_tab.refresh()
        tabs.setCurrentWidget(self.win.gallery_tab)
        tabs.setCurrentWidget(self.win.editor_tab)
        self.assertIsNone(self.win.chat_tab.browser)
        self.assertIs(self.win.editor_tab.session, self.win.session)
        self.assertIs(self.win.gallery_tab.session, self.win.session)
        self.assertIs(self.win.preview_tab.session, self.win.session)
        self.assertEqual(self.win.session.report_snapshot(), self.before)

    def test_document_edit_uses_shared_draft_and_branch_undo(self):
        child = self.before.modules[0].blocks[1].module
        block = child.blocks[0]
        address = NodeAddress(self.before.id, child.id, block.id, "text")
        self.win.tabs.setCurrentWidget(self.win.editor_tab)
        self.pump()
        editor = self.win.editor_tab._editors[address]
        editor.setPlainText("M6 approved synthetic text")
        self.pump()
        self.assertTrue(self.win.branch.drafts.pending)
        self.assertEqual(self.win.branch.drafts.draft_owner(address), "page-editor")
        self.assertEqual(self.win.session.report_snapshot(), self.before)
        with self.assertRaises(PendingDraftError):
            self.win.session.require_clean_drafts()
        self.assertTrue(self.win.editor_tab.apply_changes())
        self.pump()
        changed = self.win.session.report_snapshot()
        self.assertEqual(changed.modules[0].blocks[1].module.blocks[0].text.original,
                         "M6 approved synthetic text")
        self.assertIn("de", changed.modules[0].blocks[1].module.blocks[0].text.stale_languages)
        self.assertFalse(self.win.branch.drafts.pending)
        self.assertTrue(self.win.branch.undo())
        self.pump()
        self.assertEqual(self.win.session.report_snapshot(), self.before)

    def test_editor_revert_keeps_canonical_table_and_text_intact(self):
        table = self.before.modules[0].blocks[3]
        parent = self.before.modules[0]
        address = NodeAddress(self.before.id, parent.id, table.id, "table")
        self.win.tabs.setCurrentWidget(self.win.editor_tab)
        self.pump()
        cell = self.win.editor_tab._tables[address].item(0, 0)
        cell.setText("Not yet approved")
        self.pump()
        self.assertTrue(self.win.branch.drafts.pending)
        self.assertTrue(self.win.editor_tab.revert_changes())
        self.pump()
        self.assertFalse(self.win.branch.drafts.pending)
        self.assertEqual(self.win.session.report_snapshot(), self.before)

    def test_report_info_dialog_commits_and_pending_editor_draft_blocks_it(self):
        from reportmaker.v2.report_info import ReportInfoDialog
        from PySide6.QtGui import QAction
        action = self.win.findChild(QAction, "v2ReportInfo")
        self.assertIsNotNone(action)
        dialog = ReportInfoDialog(self.win.session, self.win.gateway, self.win)
        dialog.fields["title"].setText("Synthetic report title M6")
        dialog.types["supervision"].setChecked(True)
        dialog.types["ordinary_maintenance"].setChecked(True)
        dialog.fields["output_languages"].setText("en, de, es-MX")
        self.assertTrue(dialog.apply_changes())
        report = self.win.session.report_snapshot()
        self.assertEqual(report.title.original, "Synthetic report title M6")
        self.assertEqual(report.intervention_types,
                         ["supervision", "ordinary_maintenance"])
        self.assertEqual(report.output_languages, ["en", "de", "es-MX"])
        self.assertTrue(self.win.branch.undo())
        self.assertEqual(self.win.session.report_snapshot(), self.before)
        child = self.before.modules[0].blocks[1].module
        addr = NodeAddress(self.before.id, child.id, child.blocks[0].id, "text")
        self.win.tabs.setCurrentWidget(self.win.editor_tab)
        self.win.editor_tab._editors[addr].setPlainText("Pending editor draft")
        self.pump()
        dialog = ReportInfoDialog(self.win.session, self.win.gateway, self.win)
        dialog.fields["title"].setText("Should not commit")
        self.assertFalse(dialog.apply_changes())
        self.assertEqual(self.win.session.report_snapshot(), self.before)
        self.assertTrue(self.win.editor_tab.revert_changes())

    def test_word_template_selection_changes_export_style_not_report(self):
        from docx import Document
        from docx.shared import Cm
        with TemporaryDirectory() as directory:
            path = Path(directory) / "synthetic-layout.docx"
            source = Document()
            section = source.sections[0]
            section.page_width, section.page_height = Cm(29.7), Cm(21)
            section.left_margin, section.right_margin = Cm(2.5), Cm(2.5)
            section.top_margin, section.bottom_margin = Cm(3.0), Cm(2.5)
            section.header.paragraphs[0].text = "Synthetic header"
            section.footer.paragraphs[0].text = "Synthetic footer"
            source.save(path)
            self.assertTrue(self.win.import_layout_template(path))
            self.assertEqual(self.win.session.report_snapshot(), self.before)
            self.assertAlmostEqual(self.win.preview_tab.style.page_width_cm, 29.7,
                                   places=1)
            self.assertEqual(self.win.preview_tab.style, self.win.editor_tab.style)
            output = Path(directory) / "styled.docx"
            self.assertTrue(self.win.preview_tab.export_to(output, "docx"))
            exported = Document(output)
            self.assertAlmostEqual(exported.sections[0].page_width.cm, 29.7,
                                   places=1)
            self.assertIn("Synthetic header",
                          exported.sections[0].header.paragraphs[0].text)
            self.assertEqual(self.win.session.report_snapshot(), self.before)

    def test_preview_and_export_refuse_pending_draft(self):
        child = self.before.modules[0].blocks[1].module
        address = NodeAddress(self.before.id, child.id, child.blocks[0].id, "text")
        self.win.tabs.setCurrentWidget(self.win.editor_tab)
        self.win.editor_tab._editors[address].setPlainText("Pending")
        self.pump()
        with TemporaryDirectory() as directory:
            target = Path(directory) / "blocked.docx"
            self.assertFalse(self.win.preview_tab.export_to(target, "docx"))
            self.assertFalse(target.exists())
        self.assertTrue(self.win.editor_tab.revert_changes())
        self.win.tabs.setCurrentWidget(self.win.preview_tab)
        self.assertIn("Continuous schematic", self.win.preview_tab.notice.text())
