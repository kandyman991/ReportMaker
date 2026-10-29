"""M6 atomic language-package regression gates; synthetic content only."""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile
import json
import os
import unittest

from docx import Document

from reportmaker.export import ExportError, ExportStyle
from reportmaker.model import LocalizedText, PhotoMark, PhotoVisual
from reportmaker.v2.language_bundle import export_language_bundle
from reportmaker.v2.localized_export import TranslationRequiredError
from reportmaker.v2.projections import NodeAddress
from test_v2_m6 import synthetic_project
from test_v2_m6_followup import translate_everything


class LanguageBundlePureTests(unittest.TestCase):
    def test_all_selected_languages_both_formats_and_detached_canonical_source(self):
        project = synthetic_project()
        translate_everything(project)
        before = deepcopy(project.report)
        with TemporaryDirectory() as d:
            path = Path(d) / "languages.zip"
            self.assertEqual(export_language_bundle(project, path), path)
            with ZipFile(path) as package:
                self.assertEqual(set(package.namelist()), {
                    "manifest.json", "report_en.docx", "report_de.docx",
                    "report_en.pdf", "report_de.pdf",
                })
                manifest = json.loads(package.read("manifest.json"))
                self.assertEqual(manifest["languages"], ["en", "de"])
                self.assertEqual(len(manifest["files"]), 4)
                self.assertIn("English", manifest["metadata_localization"])
                source = Document(BytesIO(package.read("report_en.docx")))
                german = Document(BytesIO(package.read("report_de.docx")))
                source_text = "\n".join(p.text for p in source.paragraphs)
                translated_text = "\n".join(p.text for p in german.paragraphs)
                self.assertIn("Torque checked", source_text)
                self.assertIn("DE · Torque checked", translated_text)
                self.assertNotIn("DE · Torque checked", source_text)
                self.assertTrue(package.read("report_en.pdf").startswith(b"%PDF"))
                self.assertTrue(package.read("report_de.pdf").startswith(b"%PDF"))
            self.assertEqual(project.report, before)

    def test_missing_and_stale_translation_preserve_existing_archive_and_leave_no_temp(self):
        project = synthetic_project()
        with TemporaryDirectory() as d:
            directory = Path(d)
            destination = directory / "existing.zip"
            destination.write_bytes(b"previous verified output")
            with self.assertRaisesRegex(TranslationRequiredError, "Missing"):
                export_language_bundle(project, destination)
            self.assertEqual(destination.read_bytes(), b"previous verified output")
            translate_everything(project)
            project.report.modules[0].blocks[0].text.update_original("Changed since translation")
            with self.assertRaisesRegex(TranslationRequiredError, "Outdated"):
                export_language_bundle(project, destination, formats=("docx",))
            self.assertEqual(destination.read_bytes(), b"previous verified output")
            self.assertEqual(sorted(p.name for p in directory.iterdir()), ["existing.zip"])

    def test_second_language_render_failure_does_not_replace_any_zip(self):
        from reportmaker.v2 import language_bundle
        project = synthetic_project()
        translate_everything(project)
        with TemporaryDirectory() as d:
            directory = Path(d)
            destination = directory / "customer.zip"
            destination.write_bytes(b"original archive")
            real_export = language_bundle.export_project
            counter = [0]
            def second_fails(*args, **kwargs):
                counter[0] += 1
                if counter[0] == 2:
                    raise ExportError("Simulated printer backend failure")
                return real_export(*args, **kwargs)
            with patch.object(language_bundle, "export_project", side_effect=second_fails):
                with self.assertRaisesRegex(ExportError, "Simulated"):
                    export_language_bundle(project, destination, formats=("docx",))
            self.assertEqual(counter[0], 2)
            self.assertEqual(destination.read_bytes(), b"original archive")
            self.assertEqual(sorted(p.name for p in directory.iterdir()), ["customer.zip"])

    def test_subset_language_formats_and_sanitized_arbitrary_tag(self):
        project = synthetic_project()
        project.report.output_languages.append("fr-CA")
        translate_everything(project, "fr-CA")
        with TemporaryDirectory() as d:
            destination = Path(d) / "fr.zip"
            export_language_bundle(
                project, destination, languages=["fr-ca"], formats=("docx",))
            with ZipFile(destination) as package:
                self.assertEqual(set(package.namelist()),
                                 {"manifest.json", "report_fr-CA.docx"})
                self.assertIn("fr-CA", json.loads(package.read("manifest.json"))["languages"])

    def test_invalid_selection_or_archive_destination_fails_closed(self):
        project = synthetic_project()
        with TemporaryDirectory() as d:
            target = Path(d) / "target.zip"
            target.write_bytes(b"kept")
            invalid = [
                {"languages": ["de", "DE"]},
                {"languages": ["../secret"]},
                {"languages": ["fr-CA"]},
                {"languages": []},
                {"languages": ["en"] * 33},
                {"formats": ("docx", "docx")},
                {"formats": ("html",)},
                {"formats": ()},
            ]
            for options in invalid:
                with self.subTest(options=options):
                    with self.assertRaises(ValueError):
                        export_language_bundle(project, target, **options)
                    self.assertEqual(target.read_bytes(), b"kept")
            with self.assertRaises(ExportError):
                export_language_bundle(project, Path(d) / "invalid.docx")

    def test_export_style_applies_to_every_document_without_mutating_project(self):
        project = synthetic_project()
        translate_everything(project)
        before = deepcopy(project)
        brand = ExportStyle(company_name="Fake Logo-free service",
                            footer_note="Fabricated footer")
        with TemporaryDirectory() as d:
            target = Path(d) / "package.zip"
            export_language_bundle(project, target, formats=("docx",), style=brand)
            with ZipFile(target) as archive:
                for name in ("report_en.docx", "report_de.docx"):
                    document = Document(BytesIO(archive.read(name)))
                    self.assertIn("Fake Logo-free service",
                                  document.sections[0].header.paragraphs[0].text)
                    self.assertIn("Fabricated footer",
                                  document.sections[0].footer.paragraphs[0].text)
        self.assertEqual(project, before)

    def test_annotation_text_is_localized_in_package_with_explicit_renderer(self):
        from PySide6.QtWidgets import QApplication
        from reportmaker.v2.visual_export import annotated_png
        self._app = QApplication.instance() or QApplication([])
        project = synthetic_project()
        photo = project.report.modules[0].blocks[2]
        photo.visual = PhotoVisual(marks=[
            PhotoMark("text", .1, .2, .8, .4,
                      text=LocalizedText("en", "Synthetic annotation"))])
        translate_everything(project)
        before = deepcopy(project)
        with TemporaryDirectory() as d:
            target = Path(d) / "annotated.zip"
            export_language_bundle(project, target, formats=("pdf",),
                                   visual_renderer=annotated_png)
            with ZipFile(target) as package:
                self.assertTrue(package.read("report_en.pdf").startswith(b"%PDF"))
                self.assertTrue(package.read("report_de.pdf").startswith(b"%PDF"))
        self.assertEqual(project, before)


@unittest.skipUnless(os.environ.get("QT_QPA_PLATFORM") == "offscreen",
                     "Windows hosted offscreen Qt gate")
class LanguageBundleQtTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        from reportmaker.v2.app import V2MainWindow
        self.project = synthetic_project()
        translate_everything(self.project)
        self.win = V2MainWindow(project=self.project, load_website=False,
                                ephemeral_browser=True)
        self.win.show()
        self.app.processEvents()

    def tearDown(self):
        self.win._confirm_transition = lambda: False
        self.win.close()
        self.app.processEvents()

    def test_preview_package_controls_and_selected_format(self):
        from PySide6.QtWidgets import QPushButton
        preview = self.win.preview_tab
        self.assertIsNotNone(preview.findChild(QPushButton, "v2ExportLanguageBundle"))
        self.assertEqual(preview.bundle_formats.count(), 3)
        preview.bundle_formats.setCurrentIndex(1)
        with TemporaryDirectory() as d:
            target = Path(d) / "all-languages.zip"
            self.assertTrue(preview.export_bundle_to(target))
            with ZipFile(target) as package:
                self.assertEqual(set(package.namelist()),
                                 {"manifest.json", "report_en.docx", "report_de.docx"})

    def test_pending_editor_draft_blocks_package_creation_and_overwrite(self):
        report = self.win.session.report_snapshot()
        child = report.modules[0].blocks[1].module
        address = NodeAddress(report.id, child.id, child.blocks[0].id, "text")
        self.win.tabs.setCurrentWidget(self.win.editor_tab)
        self.app.processEvents()
        self.win.editor_tab._editors[address].setPlainText("Unapproved synthetic draft")
        self.app.processEvents()
        with TemporaryDirectory() as d:
            target = Path(d) / "existing.zip"
            target.write_bytes(b"existing bundle")
            self.assertFalse(self.win.preview_tab.export_bundle_to(target))
            self.assertEqual(target.read_bytes(), b"existing bundle")
        self.assertTrue(self.win.editor_tab.revert_changes())
