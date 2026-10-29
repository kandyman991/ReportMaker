"""M6.6 initial Word/table/paper fidelity gates; synthetic data only."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import os
import unittest

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from pypdf import PdfReader

from reportmaker.export import ExportStyle, export_project
from reportmaker.model import LocalizedText, TableBlock
from test_m4_export import synthetic_project


class WordPrintGeometryTests(unittest.TestCase):
    def setUp(self):
        self.project = synthetic_project()
        self.original = deepcopy(self.project)
        self.style = ExportStyle(
            page_width_cm=18, page_height_cm=25,
            left_margin_cm=1.8, right_margin_cm=2.2,
            top_margin_cm=2.1, bottom_margin_cm=1.8,
            normal_font="Calibri", normal_size_pt=12,
            heading_font="Calibri", heading_size_pt=19,
        )

    def test_word_info_and_content_tables_have_fixed_printable_grid(self):
        with TemporaryDirectory() as directory:
            target = Path(directory) / "word.docx"
            export_project(self.project, target, format="docx", style=self.style)
            doc = Document(target)
            self.assertAlmostEqual(doc.sections[0].page_width / Cm(1), 18, places=2)
            self.assertAlmostEqual(doc.sections[0].left_margin / Cm(1), 1.8, places=2)
            self.assertEqual(doc.styles["Normal"].font.size.pt, 12)
            self.assertEqual(doc.styles["Heading 1"].font.size.pt, 19)
            expected = Cm(14).twips
            self.assertGreaterEqual(len(doc.tables), 2)
            for table in doc.tables:
                with self.subTest(columns=len(table.columns)):
                    self.assertFalse(table.autofit)
                    width = table._tbl.tblPr.xpath("./w:tblW")
                    self.assertEqual(len(width), 1)
                    self.assertEqual(width[0].get(qn("w:type")), "dxa")
                    self.assertLessEqual(abs(int(width[0].get(qn("w:w"))) - expected), 2)
                    columns = sum(col.width.twips for col in table.columns)
                    self.assertLessEqual(abs(columns - expected), len(table.columns)+1)
        self.assertEqual(self.project, self.original)

    def test_wide_and_long_tables_remain_editable_and_pdf_flows_on_real_pages(self):
        # The first row is DATA, not an invented repeating header.
        words = [[LocalizedText("it", "R%02d" % i) for i in range(9)]
                 for _ in range(75)]
        self.project.report.modules[0].blocks.append(TableBlock(words))
        with TemporaryDirectory() as directory:
            word = Path(directory) / "many.docx"
            pdf = Path(directory) / "many.pdf"
            export_project(self.project, word, format="docx", style=self.style)
            export_project(self.project, pdf, format="pdf", style=self.style)
            doc = Document(word)
            table = doc.tables[-1]
            self.assertEqual(len(table.columns), 9)
            self.assertEqual(len(table.rows), 75)
            self.assertFalse(table.autofit)
            self.assertEqual(table.cell(0, 0).text, "R00")
            pages = PdfReader(pdf).pages
            self.assertGreaterEqual(len(pages), 2)
            self.assertAlmostEqual(float(pages[0].mediabox.width), Cm(18).pt, delta=1)
            self.assertAlmostEqual(float(pages[0].mediabox.height), Cm(25).pt, delta=1)


@unittest.skipUnless(os.environ.get("QT_QPA_PLATFORM") == "offscreen",
                     "Offscreen Qt geometry; real Windows DPI acceptance separate")
class NaturalPageFidelityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication
        cls.app = QApplication.instance() or QApplication([])

    def test_paper_width_and_margins_scale_together_without_fixed_500px_minimum(self):
        from reportmaker.v2.editor_tab import NaturalPagePaper
        style = ExportStyle(page_width_cm=18, left_margin_cm=2.5,
                            top_margin_cm=3, right_margin_cm=1.5,
                            bottom_margin_cm=2)
        paper = NaturalPagePaper(style)
        try:
            for width in (650, 310):
                paper.resize(width, 250)
                paper.show()
                self.app.processEvents()
                self.assertEqual(paper.width(), width)
                left, top, right, bottom = paper.flow.getContentsMargins()
                self.assertAlmostEqual(left, width * 2.5/18, delta=2)
                self.assertAlmostEqual(top, width * 3/18, delta=2)
                self.assertAlmostEqual(right, width * 1.5/18, delta=2)
                self.assertAlmostEqual(bottom, width * 2/18, delta=2)
                self.assertLess(left + right, width)
            self.assertEqual(paper.maximumWidth(), 720)
        finally:
            paper.close()

    def test_page_editor_applies_template_font_without_mutating_canonical_rich_text(self):
        from reportmaker.v2.app import V2MainWindow
        from reportmaker.v2.projections import NodeAddress
        project = synthetic_project()
        original = deepcopy(project.report)
        style = ExportStyle(normal_font="DejaVu Sans", normal_size_pt=13,
                            heading_font="DejaVu Sans", heading_size_pt=21)
        window = V2MainWindow(project=project, load_website=False,
                              ephemeral_browser=True)
        try:
            window.editor_tab.set_style(style)
            for _ in range(6):
                self.app.processEvents()
            child = original.modules[0].blocks[1].module
            paragraph = child.blocks[0]
            heading = child.blocks[-1]
            para_addr = NodeAddress(original.id, child.id, paragraph.id, "text")
            head_addr = NodeAddress(original.id, child.id, heading.id, "text")
            self.assertEqual(
                window.editor_tab._editors[para_addr].font().pointSizeF(), 13)
            self.assertEqual(
                window.editor_tab._editors[head_addr].font().pointSizeF(), 21)
            self.assertTrue(window.editor_tab._editors[head_addr].font().bold())
            self.assertEqual(window.session.report_snapshot(), original)
            self.assertFalse(window.session._editor.can_undo)
        finally:
            window._confirm_transition = lambda: False
            window.close()
            self.app.processEvents()

    def test_table_grows_on_live_cell_edit_with_inner_scrollbars_disabled(self):
        from PySide6.QtCore import Qt
        from PySide6.QtWidgets import QTableWidgetItem
        from reportmaker.v2.editor_tab import NaturalPageTable
        table = NaturalPageTable(1, 1)
        try:
            table.resize(310, 90)
            table.setItem(0, 0, QTableWidgetItem("short"))
            table.show()
            self.app.processEvents()
            before = table.sizeHint().height()
            table.item(0, 0).setText("Long synthetic table cell text " * 60)
            self.app.processEvents()
            self.assertGreater(table.sizeHint().height(), before)
            self.assertEqual(table.verticalScrollBarPolicy(),
                             Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
            self.assertEqual(table.horizontalScrollBarPolicy(),
                             Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        finally:
            table.close()


if __name__ == "__main__":
    unittest.main()
