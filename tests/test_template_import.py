"""Synthetic Word-template import regression: no real customer samples or internet."""
from __future__ import annotations

from dataclasses import replace
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED
import unittest

from docx import Document
from docx.shared import Cm, Pt
from PIL import Image
from pypdf import PdfReader

from reportmaker.export import ExportStyle, export_project
from reportmaker.template_import import (
    DocumentTemplate, TemplateError, import_word_template, load_template,
    save_template, validate_template,
)
from test_m4_export import synthetic_project


def create_sample(filename: Path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(18)
    section.page_height = Cm(25)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.9)
    section.top_margin = Cm(2.7)
    section.bottom_margin = Cm(2.2)
    section.header_distance = Cm(.8)
    section.footer_distance = Cm(.7)
    doc.styles["Normal"].font.name = "Cambria"
    doc.styles["Normal"].font.size = Pt(12)
    doc.styles["Heading 1"].font.name = "Cambria"
    doc.styles["Heading 1"].font.size = Pt(17)
    header = section.header.paragraphs[0]
    header.add_run("SYNTHETIC COMPANY")
    stream = BytesIO()
    Image.new("RGB", (40, 20), (120, 50, 20)).save(stream, "PNG")
    header.add_run().add_picture(BytesIO(stream.getvalue()), width=Cm(.9))
    section.footer.paragraphs[0].text = "Lab use only"
    doc.add_paragraph("Untrusted source document body; must not enter report")
    doc.save(filename)


def mutate_package(path, edit):
    with ZipFile(path) as z:
        parts = [(x.filename, z.read(x.filename)) for x in z.infolist()]
    rebuilt = edit(parts)
    with ZipFile(path, "w", ZIP_DEFLATED) as z:
        for name, data in rebuilt:
            z.writestr(name, data)


class TemplateImportTests(unittest.TestCase):
    def test_supported_layout_roundtrip_and_export_parity(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "sample.docx"
            create_sample(source)
            original = source.read_bytes()
            converted = import_word_template(source)
            template = converted.template
            self.assertEqual(template.page_width_cm, 18)
            self.assertEqual(template.page_height_cm, 25)
            self.assertAlmostEqual(template.left_margin_cm, 1.7, places=2)
            self.assertIn("SYNTHETIC COMPANY", template.header_text)
            self.assertEqual(template.footer_text, "Lab use only")
            self.assertEqual(template.normal_font, "Cambria")
            self.assertEqual(template.normal_size_pt, 12)
            self.assertEqual(template.heading_font, "Cambria")
            self.assertEqual(template.heading_size_pt, 17)
            self.assertTrue(template.logo)
            self.assertTrue(any("body" in w.lower() for w in converted.warnings))
            stored = root / "layout.rmtemplate"
            save_template(template, stored)
            self.assertEqual(template, load_template(stored))
            profile = ExportStyle(
                template.header_text, template.footer_text, template.logo,
                template.page_width_cm, template.page_height_cm,
                template.left_margin_cm, template.right_margin_cm,
                template.top_margin_cm, template.bottom_margin_cm,
                template.header_distance_cm, template.footer_distance_cm,
                template.normal_font, template.normal_size_pt,
                template.heading_font, template.heading_size_pt)
            project = synthetic_project()
            word, pdf = root / "output.docx", root / "output.pdf"
            export_project(project, word, format="docx", style=profile)
            export_project(project, pdf, format="pdf", style=profile)
            self.assertEqual(source.read_bytes(), original)
            generated = Document(word)
            sec = generated.sections[0]
            self.assertAlmostEqual(sec.page_width / Cm(1), 18, places=2)
            self.assertAlmostEqual(sec.page_height / Cm(1), 25, places=2)
            self.assertAlmostEqual(sec.left_margin / Cm(1), 1.7, places=2)
            self.assertEqual(generated.styles["Normal"].font.name, "Cambria")
            self.assertEqual(generated.styles["Normal"].font.size.pt, 12)
            self.assertEqual(generated.styles["Heading 1"].font.size.pt, 17)
            self.assertIn("SYNTHETIC COMPANY", sec.header.paragraphs[0].text)
            self.assertIn("Lab use only", sec.footer.paragraphs[0].text)
            self.assertNotIn("Untrusted source document body", "\n".join(p.text for p in generated.paragraphs))
            pdf_pages = PdfReader(pdf).pages
            self.assertGreaterEqual(len(pdf_pages), 1)
            self.assertAlmostEqual(float(pdf_pages[0].mediabox.width), 18*72/2.54, delta=2)
            self.assertAlmostEqual(float(pdf_pages[0].mediabox.height), 25*72/2.54, delta=2)
            self.assertIn("Lab use only", pdf_pages[0].extract_text())

    def test_dotx_main_type_is_normalized_only_in_memory(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / "sample.dotx"
            create_sample(path)
            def make_dotx(parts):
                return [(name, data.replace(
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
                ) if name == "[Content_Types].xml" else data)
                        for name, data in parts]
            mutate_package(path, make_dotx)
            original = path.read_bytes()
            loaded = import_word_template(path)
            self.assertEqual(loaded.template.page_height_cm, 25)
            self.assertEqual(path.read_bytes(), original)

    def test_external_relationship_rejected(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "bad.docx"
            create_sample(path)
            def link(parts):
                return [(name, payload.replace(
                    b"</Relationships>",
                    b'<Relationship Id="external" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="https://example.invalid/track" TargetMode="External"/></Relationships>'
                ) if name == "word/_rels/document.xml.rels" else payload)
                        for name, payload in parts]
            mutate_package(path, link)
            with self.assertRaisesRegex(TemplateError, "External"):
                import_word_template(path)

    def test_macro_part_zip_traversal_and_xml_dtd_rejected(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / "bad.docx"
            create_sample(path)
            mutate_package(path, lambda parts: parts + [
                ("word/vbaProject.bin", b"not a real macro")])
            with self.assertRaisesRegex(TemplateError, "Active"):
                import_word_template(path)
            create_sample(path)
            mutate_package(path, lambda parts: parts + [("../escape", b"no")])
            with self.assertRaisesRegex(TemplateError, "path"):
                import_word_template(path)
            create_sample(path)
            mutate_package(path, lambda parts: [
                (name, b'<!DOCTYPE doc [<!ENTITY leak SYSTEM "file:///etc/passwd">]>' + data
                 if name == "word/document.xml" else data)
                for name, data in parts])
            with self.assertRaisesRegex(TemplateError, "DTD"):
                import_word_template(path)

    def test_benign_zip_directories_and_corruption_handling(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "word.docx"
            create_sample(path)
            mutate_package(path, lambda parts: parts + [("additional/", b"")])
            self.assertEqual(import_word_template(path).template.page_width_cm, 18)
            path.write_bytes(b"PK" + b"broken" * 10)
            with self.assertRaises(TemplateError):
                import_word_template(path)

    def test_corrupt_profile_and_geometry_rejected_without_overwriting(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            target = root / "layout.rmtemplate"
            base = DocumentTemplate()
            save_template(base, target)
            good = target.read_bytes()
            for update in (
                {"page_width_cm": 0},
                {"left_margin_cm": 7, "right_margin_cm": 7, "page_width_cm": 15},
                {"header_text": "\nHTML injection"},
                {"normal_size_pt": float("nan")},
                {"heading_size_pt": 120},
                {"logo": b"not an image"},
            ):
                with self.assertRaises((ValueError, TypeError)):
                    save_template(replace(base, **update), target)
                self.assertEqual(target.read_bytes(), good)
            target.write_text('{"format":"unknown","layout":{},"version":1}')
            with self.assertRaises(TemplateError):
                load_template(target)
            target.write_bytes(good + b"extra")
            with self.assertRaises(TemplateError):
                load_template(target)

    def test_qt_import_selection_and_clear_is_local_and_non_destructive(self):
        try:
            from PySide6.QtCore import QSettings
            from PySide6.QtWidgets import (
                QApplication, QFileDialog, QHBoxLayout, QMainWindow,
                QMessageBox, QVBoxLayout, QWidget,
            )
            from reportmaker.visual_editor import VisualReportEditor
        except ImportError:
            self.skipTest("Qt desktop dependencies absent")
        app = QApplication.instance() or QApplication([])
        with TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "template.docx"
            out = root / "kept.rmtemplate"
            create_sample(source)
            window = QMainWindow()
            window.settings = QSettings(str(root / "prefs.ini"), QSettings.Format.IniFormat)
            host = QWidget()
            window.setCentralWidget(host)
            columns = QHBoxLayout(host)
            a, b = QVBoxLayout(), QVBoxLayout()
            columns.addLayout(a)
            columns.addLayout(b)
            editor = VisualReportEditor(window, a, b)
            with (patch.object(QFileDialog, "getOpenFileName", return_value=(str(source), "")),
                  patch.object(QFileDialog, "getSaveFileName", return_value=(str(out), "")),
                  patch("reportmaker.visual_editor.edit_template_layout",
                        side_effect=lambda owner, template, warnings=(): template),
                  patch.object(QMessageBox, "warning",
                               side_effect=AssertionError("Unexpected template modal warning"))):
                self.assertTrue(editor.import_document_template())
            active = editor._export_style()
            self.assertEqual(active.page_width_cm, 18)
            self.assertEqual(active.heading_size_pt, 17)
            self.assertEqual(editor.page.layout_template.page_width_cm, 18)
            self.assertIn("SYNTHETIC", active.company_name)
            self.assertEqual(load_template(out).page_width_cm, 18)
            self.assertTrue(editor.clear_document_template())
            self.assertEqual(editor._export_style().page_width_cm, 21)
            self.assertTrue(out.exists())
            window.close()


if __name__ == "__main__":
    unittest.main()
