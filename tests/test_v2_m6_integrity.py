"""M6 independent-review gates H1–H5; fabricated reports only."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import os
import unittest

from reportmaker.export import export_project, prepare_export, ExportStyle
from reportmaker.model import LocalizedText
from reportmaker.rich_text import RichFormatting, TextRun, ParagraphStyle
from reportmaker.v2.commands import CommandGateway, CommandIntent
from reportmaker.v2.drafts import DraftCoordinator, TextDraft
from reportmaker.v2.localized_export import (
    TranslationRequiredError, localized_project, localized_fields,
)
from reportmaker.v2.projections import NodeAddress
from reportmaker.v2.session import ReportSession
from test_v2_m6 import synthetic_project
from test_v2_m6_followup import translate_everything


class ExportAndIntegrityRegressionTests(unittest.TestCase):
    def test_h1_shorter_multiline_unicode_translation_never_reuses_source_offsets(self):
        project = synthetic_project()
        source = project.report.modules[0].blocks[1].module.blocks[0]
        source.rich = RichFormatting(
            runs=[TextRun(0, len(source.text.original), bold=True)],
            paragraphs=[ParagraphStyle(0, align="center")],
        )
        translate_everything(project)
        source.text.set_translation("de", "Ü\nß")
        untouched = deepcopy(project.report)
        detached = localized_project(project, "de")
        result = detached.report.modules[0].blocks[1].module.blocks[0]
        self.assertEqual(result.text.original, "Ü\nß")
        self.assertEqual(result.rich, RichFormatting())
        self.assertEqual(project.report, untouched)
        self.assertEqual(source.rich.runs[0].end, len(source.text.original))
        plan = prepare_export(detached)
        self.assertTrue(any(item.text == "Ü\nß" for item in plan.items))
        with TemporaryDirectory() as d:
            destination = Path(d) / "translated.docx"
            export_project(detached, destination, format="docx")
            self.assertGreater(destination.stat().st_size, 500)

    def test_h4_blank_translation_rejected_on_write_and_even_from_loaded_project(self):
        project = synthetic_project()
        session, gateway = ReportSession(project), None
        gateway = CommandGateway(session)
        original = session.report_snapshot()
        for bad in ("", "   "):
            with self.assertRaisesRegex(ValueError, "nonblank"):
                gateway.execute(CommandIntent(
                    "set_translation", None,
                    {"key": (original.id, "report_title"), "language": "de", "text": bad},
                    session.revision, session.generation))
        self.assertEqual(session.revision, 0)
        self.assertEqual(session.report_snapshot(), original)
        translate_everything(project)
        for key, value in localized_fields(project.report):
            if key[1] in ("module_title", "caption", "cell", "text"):
                if value.original:
                    value.set_translation("de", "")
                    with self.assertRaisesRegex(TranslationRequiredError, "Blank"):
                        localized_project(project, "de")
                    break
        else:
            self.fail("Expected at least one nonempty localized block")

    def test_h4_empty_source_can_retain_intentional_target_only_content(self):
        project = synthetic_project()
        translate_everything(project)
        project.report.conclusion.blocks.append(
            __import__("reportmaker.model", fromlist=["ParagraphBlock"]).ParagraphBlock(
                LocalizedText("en", "", translations={"de": "Nur Deutsch"})))
        output = localized_project(project, "de")
        self.assertEqual(output.report.conclusion.blocks[-1].text.original, "Nur Deutsch")
        self.assertEqual(project.report.conclusion.blocks[-1].text.original, "")

    def test_h3_primary_language_switch_requires_consent_and_relabels_all_fields(self):
        project = synthetic_project()
        session = ReportSession(project)
        gateway = CommandGateway(session)
        before = session.report_snapshot()
        args = {
            "title": before.title.original,
            "created_date": before.created_date,
            "start_date": before.start_date,
            "end_date": before.end_date,
            "ticket": before.ticket,
            "technicians": before.technicians,
            "intervention_types": before.intervention_types,
            "other_intervention_detail": before.other_intervention_detail,
            "customer_snapshot": before.customer_snapshot,
            "customer_id": before.customer_id,
            "primary_language": "fr-CA",
            "output_languages": ["en", "de", "fr-CA"],
        }
        def intent():
            return CommandIntent("set_report_metadata", None, args,
                                 session.revision, session.generation)
        with self.assertRaisesRegex(ValueError, "explicit confirmation"):
            gateway.execute(intent())
        self.assertEqual(session.report_snapshot(), before)
        self.assertEqual(session.revision, 0)
        args["relabel_source_language"] = True
        gateway.execute(intent())
        after = session.report_snapshot()
        self.assertEqual(after.primary_language, "fr-CA")
        self.assertEqual(after.title.original, before.title.original)
        self.assertEqual(after.modules, [
            __import__("copy").deepcopy(module) for module in after.modules])
        for _key, text in localized_fields(after):
            self.assertEqual(text.original_language, "fr-CA")
            self.assertTrue(set(text.translations).issubset(text.stale_languages))
        self.assertTrue(gateway.undo(session.generation, session.revision))
        self.assertEqual(session.report_snapshot(), before)
        self.assertTrue(gateway.redo(session.generation, session.revision))
        self.assertEqual(session.report_snapshot(), after)
        collision = dict(args, primary_language="de")
        with self.assertRaises(ValueError):
            gateway.execute(CommandIntent("set_report_metadata", None, collision,
                                          session.revision, session.generation))
        self.assertEqual(session.report_snapshot(), after)

    def test_mixed_content_translation_fields_follow_document_order(self):
        report = synthetic_project().report
        first, second = report.modules
        keys = [key for key, _value in localized_fields(report)]
        heading_key = (first.blocks[0].id, "text")
        child_key = (first.blocks[1].module.id, "module_title")
        photo_key = (first.blocks[2].id, "caption")
        table_key = (first.blocks[3].id, "cell", 0, 0)
        self.assertLess(keys.index(heading_key), keys.index(child_key))
        self.assertLess(keys.index(child_key), keys.index(photo_key))
        self.assertLess(keys.index(photo_key), keys.index(table_key))
        self.assertLess(keys.index((second.id, "module_title")),
                        keys.index((report.conclusion.id, "module_title")))

    def test_h2_multiple_drafts_commit_one_undo_and_failed_batch_retains_all(self):
        project = synthetic_project()
        session = ReportSession(project)
        gateway = CommandGateway(session)
        drafts = DraftCoordinator(session, gateway)
        port = drafts.for_editor("page-editor")
        before = session.report_snapshot()
        first = before.modules[0].blocks[0]
        child = before.modules[0].blocks[1].module
        second = child.blocks[0]
        a = NodeAddress(before.id, before.modules[0].id, first.id, "text")
        b = NodeAddress(before.id, child.id, second.id, "text")
        port.stage_text(a, "Synthetic heading replaced")
        port.stage_text(b, "Synthetic paragraph replaced")
        self.assertTrue(port.owns(a) and port.owns(b))
        # Simulate a corrupted queued rich range in the second draft.
        old = drafts._drafts[b]
        drafts._drafts[b] = TextDraft(b, old.original, old.value, old.generation,
                                      RichFormatting(runs=[TextRun(0, 10000)]))
        with self.assertRaises(ValueError):
            port.apply_batch((a, b))
        self.assertTrue(port.owns(a) and port.owns(b))
        self.assertEqual(session.report_snapshot(), before)
        self.assertEqual(session.revision, 0)
        drafts._drafts[b] = old
        self.assertIsNotNone(port.apply_batch((a, b)))
        self.assertEqual(session.revision, 1)
        self.assertFalse(port.owns(a) or port.owns(b))
        changed = session.report_snapshot()
        self.assertEqual(changed.modules[0].blocks[0].text.original,
                         "Synthetic heading replaced")
        self.assertEqual(changed.modules[0].blocks[1].module.blocks[0].text.original,
                         "Synthetic paragraph replaced")
        self.assertTrue(gateway.undo(session.generation, session.revision))
        self.assertEqual(session.report_snapshot(), before)
        self.assertTrue(gateway.redo(session.generation, session.revision))
        self.assertEqual(session.report_snapshot(), changed)


@unittest.skipUnless(os.environ.get("QT_QPA_PLATFORM") == "offscreen",
                     "Hosted Windows offscreen Qt regression")
class QtIntegrityRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        from reportmaker.v2.app import V2MainWindow
        self.win = V2MainWindow(project=synthetic_project(),
                                load_website=False, ephemeral_browser=True)
        self.win.show()
        self.app.processEvents()

    def tearDown(self):
        self.win._confirm_transition = lambda: False
        self.win.close()
        self.app.processEvents()

    def test_h2_page_apply_two_fields_uses_one_revision_and_one_undo(self):
        before = self.win.session.report_snapshot()
        parent = before.modules[0]
        child = parent.blocks[1].module
        a = NodeAddress(before.id, parent.id, parent.blocks[0].id, "text")
        b = NodeAddress(before.id, child.id, child.blocks[0].id, "text")
        self.win.tabs.setCurrentWidget(self.win.editor_tab)
        self.app.processEvents()
        self.win.editor_tab._editors[a].setPlainText("New heading")
        self.win.editor_tab._editors[b].setPlainText("New paragraph")
        self.app.processEvents()
        self.assertTrue(self.win.editor_tab.apply_changes())
        self.assertEqual(self.win.session.revision, 1)
        self.assertTrue(self.win.branch.undo())
        self.assertEqual(self.win.session.report_snapshot(), before)

    def test_linked_reference_identity_edits_require_explicit_unlink(self):
        from reportmaker.reference import CustomerSnapshot
        from reportmaker.v2.report_info import ReportInfoDialog

        original = self.win.session.report_snapshot()
        dialog = ReportInfoDialog(self.win.session, self.win.gateway, self.win)
        linked = CustomerSnapshot(
            customer_id="fictional-db-customer", customer_name="Customer A",
            site_id="fictional-db-site", site_name="Site A",
            machine_id="fictional-db-machine", machine_name="Machine A",
            machine_serial="FAKE-001", customer_address="Fictional street")
        dialog._snapshot = linked  # Simulate a successful database selection.
        for key in ("customer_name", "site_name", "machine_name", "machine_serial"):
            dialog.fields[key].setText(getattr(linked, key))
        self.assertEqual(dialog._build_args()["customer_id"], linked.customer_id)
        for key in ("customer_name", "site_name", "machine_name", "machine_serial"):
            with self.subTest(field=key):
                dialog.fields[key].setText("Different identity")
                with self.assertRaisesRegex(ValueError, "one-off manual"):
                    dialog._build_args()
                self.assertFalse(dialog.apply_changes())
                self.assertEqual(self.win.session.revision, 0)
                self.assertEqual(self.win.session.report_snapshot(), original)
                dialog.fields[key].setText(getattr(linked, key))

        dialog.fields["customer_name"].setText("Different customer")
        self.assertTrue(dialog.use_manual())
        args = dialog._build_args()
        self.assertIsNone(args["customer_id"])
        self.assertIsNone(args["customer_snapshot"].site_id)
        self.assertIsNone(args["customer_snapshot"].machine_id)
        self.assertEqual(args["customer_snapshot"].customer_name, "Different customer")
        self.assertEqual(args["customer_snapshot"].customer_address, "Fictional street")
        self.assertTrue(dialog.apply_changes())
        saved = self.win.session.report_snapshot()
        self.assertIsNone(saved.customer_id)
        self.assertEqual(saved.customer_snapshot.customer_name, "Different customer")
        self.assertTrue(self.win.branch.undo())
        self.assertEqual(self.win.session.report_snapshot(), original)

    def test_h5_new_and_open_project_reset_another_customers_branding(self):
        from reportmaker.export import ExportStyle
        from docx import Document
        style = ExportStyle(company_name="Synthetic customer A brand",
                            footer_note="Customer A private footer")
        with TemporaryDirectory() as d:
            old_path = Path(d) / "customer-a.rmproj"
            self.assertTrue(self.win.save_project(old_path))
            self.win.editor_tab.set_style(style)
            self.win.preview_tab.set_style(style)
            self.assertEqual(self.win.preview_tab.style, style)
            self.assertTrue(self.win.new_project("Fictional customer B"))
            self.assertEqual(self.win.editor_tab.style, ExportStyle())
            self.assertEqual(self.win.preview_tab.style, ExportStyle())
            self.win.editor_tab.set_style(style)
            self.win.preview_tab.set_style(style)
            with patch.object(self.win, "_confirm_transition", return_value=True):
                self.assertTrue(self.win.open_project(old_path))
            self.assertEqual(self.win.preview_tab.style, ExportStyle())
            target = Path(d) / "customer-a-export.docx"
            self.assertTrue(self.win.preview_tab.export_to(target, "docx"))
            word = Document(target)
            self.assertNotIn("Synthetic customer A brand",
                             "\n".join(p.text for p in word.paragraphs))
            self.assertNotIn("Customer A private footer",
                             "\n".join(p.text for s in word.sections
                                       for p in s.footer.paragraphs))
