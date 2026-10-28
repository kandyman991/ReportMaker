"""Synthetic M4 branding/layout regressions, no user assets or cloud access."""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document
from PIL import Image
from pypdf import PdfReader
from pypdf.generic import ContentStream
from reportlab.lib.units import cm

from reportmaker.export import (
    ExportError, ExportStyle, PHOTO_MAX_HEIGHT_CM, PHOTO_MAX_WIDTH_CM,
    _logo_dimensions, _photo_dimensions, export_project,
)
from reportmaker.project_io import save_project
from test_m4_export import synthetic_project


def synthetic_logo(size=(240, 80)) -> bytes:
    image = Image.new("RGBA", size, (25, 90, 140, 255))
    stream = BytesIO()
    image.save(stream, "PNG")
    return stream.getvalue()


def first_logo_placement(page, reader):
    """Read the actual PDF image transform, not only the source dimensions."""
    operations = ContentStream(page.get_contents(), reader).operations
    for i, (args, opcode) in enumerate(operations):
        if opcode == b"Do" and str(args[0]).startswith("/FormXob"):
            transform, kind = operations[i - 1]
            if kind == b"cm":
                return tuple(float(value) for value in transform)
    raise AssertionError("No placed header logo image found in PDF contents")


class M4BrandingTests(unittest.TestCase):
    def test_word_repeated_branding_footer_field_and_keep_caption(self):
        project = synthetic_project()
        before = deepcopy(project)
        style = ExportStyle("Sample Service GmbH", "Synthetic test • internal", synthetic_logo())
        with TemporaryDirectory() as directory:
            source = Path(directory) / "source.rmproj"
            save_project(project, source)
            initial = source.read_bytes()
            dest = Path(directory) / "branded.docx"
            export_project(project, dest, format="docx", style=style)
            doc = Document(dest)
            section = doc.sections[0]
            self.assertIn("Sample Service GmbH", section.header.paragraphs[0].text)
            self.assertIn("Synthetic test", section.footer.paragraphs[0].text)
            self.assertEqual(len(section.header.paragraphs[0].runs[0]._r.xpath(".//a:blip")), 1)
            self.assertGreaterEqual(len(doc.inline_shapes), 1)
            self.assertIn("Foto dimostrativa", [p.text for p in doc.paragraphs])
            picture = next(p for p in doc.paragraphs if p._p.xpath(".//a:blip"))
            self.assertTrue(picture.paragraph_format.keep_with_next)
            with __import__("zipfile").ZipFile(dest) as archive:
                footer_file = next(n for n in archive.namelist() if n.startswith("word/footer") and n.endswith(".xml"))
                xml = archive.read(footer_file).decode()
                self.assertIn("PAGE", xml)
            self.assertEqual(project, before)
            self.assertEqual(source.read_bytes(), initial)

    def test_pdf_branded_header_and_footer_on_every_page(self):
        project = synthetic_project()
        style = ExportStyle("Sample Service GmbH", "Synthetic test footer", synthetic_logo())
        with TemporaryDirectory() as directory:
            target = Path(directory) / "branded.pdf"
            export_project(project, target, format="pdf", style=style)
            reader = PdfReader(target)
            self.assertGreaterEqual(len(reader.pages), 2)
            for i, page in enumerate(reader.pages, 1):
                extracted = page.extract_text()
                self.assertIn("Sample Service GmbH", extracted)
                self.assertIn("Synthetic test footer", extracted)
                self.assertIn("Page " + str(i), extracted)
                images = page["/Resources"].get("/XObject")
                self.assertIsNotNone(images)
                w, _, _, h, x, y = first_logo_placement(page, reader)
                expected_w, expected_h = _logo_dimensions(style.logo)
                self.assertAlmostEqual(w, expected_w*cm, delta=.1)
                self.assertAlmostEqual(h, expected_h*cm, delta=.1)
                self.assertAlmostEqual(x, style.left_margin_cm*cm, delta=.1)
                self.assertGreaterEqual(y + .1,
                                        (style.page_height_cm - style.top_margin_cm)*cm)

    def test_oversized_pdf_logo_scales_to_header_without_overlap(self):
        # The source is taller than the available header. It must be present
        # on every page at its correct aspect ratio and above the body frame.
        project = synthetic_project()
        logo = synthetic_logo((900, 700))
        style = ExportStyle("Synthetic letterhead", "Synthetic footer", logo,
                            top_margin_cm=1.8, header_distance_cm=1.1)
        room = style.top_margin_cm - style.header_distance_cm - .2
        natural_w, natural_h = _logo_dimensions(logo)
        expected_w, expected_h = _logo_dimensions(logo, max_height_cm=room)
        self.assertLess(expected_h, natural_h)
        self.assertAlmostEqual(expected_w / expected_h, natural_w / natural_h)
        with TemporaryDirectory() as directory:
            target = Path(directory) / "large-header-logo.pdf"
            export_project(project, target, format="pdf", style=style)
            reader = PdfReader(target)
            self.assertGreaterEqual(len(reader.pages), 2)
            for i, page in enumerate(reader.pages, 1):
                self.assertIn("Synthetic letterhead", page.extract_text())
                self.assertIn("Synthetic footer", page.extract_text())
                self.assertIn("Page " + str(i), page.extract_text())
                w, skew_x, skew_y, h, x, y = first_logo_placement(page, reader)
                self.assertAlmostEqual(w, expected_w*cm, delta=.1)
                self.assertAlmostEqual(h, expected_h*cm, delta=.1)
                self.assertAlmostEqual(skew_x, 0)
                self.assertAlmostEqual(skew_y, 0)
                self.assertAlmostEqual(x, style.left_margin_cm*cm, delta=.1)
                self.assertAlmostEqual(y + h,
                    (style.page_height_cm - style.header_distance_cm - .2)*cm,
                    delta=.1)
                self.assertGreaterEqual(y + .1,
                    (style.page_height_cm - style.top_margin_cm)*cm)

    def test_very_tight_but_positive_header_can_export_logo(self):
        project = synthetic_project()
        style = ExportStyle(logo=synthetic_logo(),
                            top_margin_cm=1.6, header_distance_cm=1.25)
        with TemporaryDirectory() as directory:
            path = Path(directory) / "tight-header.pdf"
            export_project(project, path, format="pdf", style=style)
            reader = PdfReader(path)
            self.assertGreaterEqual(len(reader.pages), 2)
            for page in reader.pages:
                _, _, _, height, _, y = first_logo_placement(page, reader)
                self.assertAlmostEqual(height,
                    (style.top_margin_cm-style.header_distance_cm-.2)*cm,
                    delta=.1)
                self.assertGreaterEqual(y + .1,
                    (style.page_height_cm-style.top_margin_cm)*cm)

    def test_never_enlarge_photos_and_fit_extreme_aspects(self):
        for width, height in ((20, 12), (4000, 100), (100, 4000)):
            stream = BytesIO()
            Image.new("RGB", (width, height), "green").save(stream, "PNG")
            out_w, out_h = _photo_dimensions(stream.getvalue())
            self.assertLessEqual(out_w, PHOTO_MAX_WIDTH_CM + .0001)
            self.assertLessEqual(out_h, PHOTO_MAX_HEIGHT_CM + .0001)
            self.assertLessEqual(out_w, width * 2.54 / 96 + .0001)
            self.assertLessEqual(out_h, height * 2.54 / 96 + .0001)

    def test_invalid_branding_does_not_clobber_output(self):
        project = synthetic_project()
        with TemporaryDirectory() as directory:
            for kind in ("docx", "pdf"):
                target = Path(directory) / ("existing." + kind)
                target.write_bytes(b"previous approved export")
                for style in (
                    ExportStyle(company_name="bad\x00company"),
                    ExportStyle(logo=b"not an image"),
                    ExportStyle(logo=b"x" * (2*1024*1024 + 1)),
                ):
                    with self.subTest(format=kind, style=str(style)[:45]):
                        with self.assertRaises(ExportError):
                            export_project(project, target, format=kind, style=style)
                        self.assertEqual(target.read_bytes(), b"previous approved export")
                self.assertFalse(list(Path(directory).glob(".reportmaker-export-*")))
            crowded = Path(directory) / "crowded.pdf"
            crowded.write_bytes(b"previous approved export")
            with self.assertRaisesRegex(ExportError, "Footer text does not fit"):
                export_project(
                    project, crowded, format="pdf",
                    style=ExportStyle(footer_note="Synthetic long footer " * 75))
            self.assertEqual(crowded.read_bytes(), b"previous approved export")
            # No geometric space at all is genuinely impossible. Validate
            # before ReportLab's onPage callback and preserve the old output.
            obstructed = Path(directory) / "logo-overlap.pdf"
            obstructed.write_bytes(b"previous approved export")
            with self.assertRaisesRegex(ExportError, "PDF header has no space") as error:
                export_project(
                    project, obstructed, format="pdf",
                    style=ExportStyle(logo=synthetic_logo(),
                                      top_margin_cm=1.6, header_distance_cm=1.55))
            self.assertIn("increase the top margin", str(error.exception))
            self.assertNotIn("handle_pageBegin", str(error.exception))
            self.assertEqual(obstructed.read_bytes(), b"previous approved export")

            # Header text cannot be silently truncated to accommodate a logo.
            crowded_header = Path(directory) / "crowded-header.pdf"
            crowded_header.write_bytes(b"previous approved export")
            with self.assertRaisesRegex(ExportError, "Header text does not fit") as error:
                export_project(
                    project, crowded_header, format="pdf",
                    style=ExportStyle(company_name="First line\nSecond line",
                                      logo=synthetic_logo(),
                                      top_margin_cm=1.6, header_distance_cm=1.0))
            self.assertNotIn("handle_pageBegin", str(error.exception))
            self.assertEqual(crowded_header.read_bytes(), b"previous approved export")
            self.assertFalse(list(Path(directory).glob(".reportmaker-export-*")))


if __name__ == "__main__":
    unittest.main()
