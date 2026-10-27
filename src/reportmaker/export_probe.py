"""Offline hosted-Windows smoke: Qt M4 export without browser/network access."""
from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from PySide6.QtCore import QSettings
from PySide6.QtGui import QAction
from PySide6.QtWidgets import QApplication, QMessageBox
from docx import Document
from pypdf import PdfReader
from PIL import Image

from .project_io import load_project
from .ui import MainWindow


def run() -> int:
    app = QApplication.instance() or QApplication([])
    with TemporaryDirectory(prefix="reportmaker-m4-") as directory:
        root = Path(directory)
        window = MainWindow(
            load_website=False,
            settings=QSettings(str(root / "settings.ini"), QSettings.Format.IniFormat),
            ephemeral_browser=True,
        )
        window.show()
        app.processEvents()
        doc = window.document
        try:
            assert window.findChild(QAction, "exportDOCXAction") is not None
            assert window.findChild(QAction, "exportPDFAction") is not None
            assert window.findChild(QAction, "exportBrandingAction") is not None
            logo = root / "company-logo.png"
            Image.new("RGBA", (250, 75), (36, 90, 140, 255)).save(logo)
            window.settings.setValue("export/company_name", "Synthetic Export Company")
            window.settings.setValue("export/footer_note", "Test-only document")
            window.settings.setValue("export/logo_path", str(logo))
            doc.info["customer"].setText("Sample Export Customer")
            doc.info["other"].setText("Synthetic inspection")
            doc.types["other"].setChecked(True)
            doc._set_info_pending()
            source = root / "project.rmproj"
            assert doc.save_to_path(source)
            original = source.read_bytes()
            doc.info["ticket"].setText("M4-CHANGED-UNSAVED")
            doc._set_info_pending()
            word = root / "report.docx"
            pdf = root / "report.pdf"
            assert doc.export_to_path(word, "docx")
            assert doc.export_to_path(pdf, "pdf")
            assert "Synthetic Export Company" in Document(word).sections[0].header.paragraphs[0].text
            assert "Test-only document" in Document(word).sections[0].footer.paragraphs[0].text
            assert "Synthetic Export Company" in PdfReader(pdf).pages[0].extract_text()
            assert "Test-only document" in PdfReader(pdf).pages[-1].extract_text()
            assert doc.path == source
            assert doc.dirty  # Export must not masquerade as project Save.
            assert source.read_bytes() == original
            assert load_project(source).report.ticket == ""
            assert "M4-CHANGED-UNSAVED" in [
                cell.text for table in Document(word).tables for row in table.rows for cell in row.cells
            ]
            assert "M4-CHANGED-UNSAVED" in "\n".join(
                page.extract_text() for page in PdfReader(pdf).pages
            )
            # A missing configured logo must fail without altering an earlier export.
            before = word.read_bytes()
            window.settings.setValue("export/logo_path", str(root / "missing.png"))
            with patch.object(QMessageBox, "warning",
                              return_value=QMessageBox.StandardButton.Ok):
                assert not doc.export_to_path(word, "docx")
            assert word.read_bytes() == before
            window.settings.setValue("export/logo_path", str(logo))
            # Invalid draft/format must not destroy an existing export.
            doc.types["other"].setChecked(False)
            doc.info["other"].setText("inconsistent detail")
            doc._set_info_pending()
            with patch.object(QMessageBox, "warning",
                              return_value=QMessageBox.StandardButton.Ok):
                assert not doc.export_to_path(word, "docx")
            assert word.read_bytes() == before
            # Reopen original source then confirm clean project state.
            assert doc.open_from_path(source)
            assert not doc.dirty
        finally:
            doc.dirty = False
            window.close()
            app.processEvents()
    return 0
