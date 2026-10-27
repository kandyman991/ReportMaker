"""Local, bounded Word layout-template adapter. Never execute, extract or fetch OOXML.

A Word package is *input*, not a document to edit in-place. We save only a
validated, versioned presentation profile (.rmtemplate), independent of report
facts and external user files. Unsupported constructs are disclosed at import.
"""
from __future__ import annotations

from base64 import b64decode, b64encode
from dataclasses import asdict, dataclass
from io import BytesIO
import json
import os
from pathlib import Path
import re
import tempfile
from zipfile import BadZipFile, ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

MAX_SOURCE = 24 * 1024 * 1024
MAX_UNPACKED = 40 * 1024 * 1024
MAX_XML = 8 * 1024 * 1024
MAX_ITEMS = 400
MAX_TEMPLATE = 3 * 1024 * 1024
MAX_LOGO = 2 * 1024 * 1024
FORMAT = "reportmaker.layout-template"
VERSION = 1

_BLOCKED = ("vbaproject", "vba", "activex", "embeddings/", "oleobject",
            "webextensions/", "attachedtemplate", "altchunk", "customui/",
            "macros/", "scripts/", "externallink")
_DANGEROUS_REL = ("oleobject", "attachedtemplate", "afchunk", "vbaproject",
                  "control", "package", "external")


class TemplateError(ValueError):
    """Untrusted or unsupported local Word template."""


@dataclass(frozen=True)
class DocumentTemplate:
    page_width_cm: float = 21.0
    page_height_cm: float = 29.7
    left_margin_cm: float = 2.1
    right_margin_cm: float = 2.1
    top_margin_cm: float = 3.0
    bottom_margin_cm: float = 2.5
    header_distance_cm: float = 0.9
    footer_distance_cm: float = 0.9
    header_text: str = ""
    footer_text: str = ""
    normal_font: str = ""
    normal_size_pt: float = 11.0
    heading_font: str = ""
    heading_size_pt: float = 16.0
    logo: bytes = b""


@dataclass(frozen=True)
class ImportResult:
    template: DocumentTemplate
    warnings: tuple[str, ...]


def validate_template(template: DocumentTemplate) -> DocumentTemplate:
    from math import isfinite
    if not isinstance(template, DocumentTemplate):
        raise TemplateError("Invalid layout template record")
    dimensions = (template.page_width_cm, template.page_height_cm)
    margins = (template.left_margin_cm, template.right_margin_cm,
               template.top_margin_cm, template.bottom_margin_cm)
    offsets = (template.header_distance_cm, template.footer_distance_cm)
    if any(type(x) not in (int, float) or not isfinite(x) for x in
           dimensions + margins + offsets + (template.normal_size_pt, template.heading_size_pt)):
        raise TemplateError("Nonfinite or invalid layout geometry")
    if not (15 <= dimensions[0] <= 43 and 15 <= dimensions[1] <= 60):
        raise TemplateError("Unsupported page dimensions")
    if any(not 0.5 <= x <= 7 for x in margins):
        raise TemplateError("Unsupported page margins")
    if (dimensions[0] - margins[0] - margins[1] < 9 or
        dimensions[1] - margins[2] - margins[3] < 9):
        raise TemplateError("Printable page area is too small")
    if (not 0.3 <= offsets[0] <= margins[2] or
        not 0.3 <= offsets[1] <= margins[3]):
        raise TemplateError("Header/footer position exceeds page margins")
    for value, limit in ((template.header_text, 120), (template.footer_text, 120),
                         (template.normal_font, 96), (template.heading_font, 96)):
        if (not isinstance(value, str) or len(value) > limit or
            any(ord(char) < 32 or ord(char) == 127 for char in value)):
            raise TemplateError("Unsupported header/footer/font text")
    for face in (template.normal_font, template.heading_font):
        if face and not re.fullmatch(r"[\w .-]{1,96}", face):
            raise TemplateError("Unsupported Word font name")
    if not 6 <= template.normal_size_pt <= 72 or not 8 <= template.heading_size_pt <= 40:
        raise TemplateError("Unsupported Word font point size")
    if not isinstance(template.logo, bytes) or len(template.logo) > MAX_LOGO:
        raise TemplateError("Embedded logo exceeds size limit")
    if template.logo:
        from .export import _image_bytes, ExportError
        try:
            _image_bytes(template.logo)
        except ExportError as exc:
            raise TemplateError("Embedded logo must be a safe local image") from exc
    return template


def _inspect_package(blob: bytes) -> tuple[bytes, tuple[str, ...]]:
    if len(blob) > MAX_SOURCE or len(blob) < 40:
        raise TemplateError("Word template size is unsupported")
    try:
        archive = ZipFile(BytesIO(blob))
    except BadZipFile as exc:
        raise TemplateError("Not a valid DOCX/DOTX package") from exc
    warnings = []
    with archive:
        infos = archive.infolist()
        names = [item.filename for item in infos]
        if (not 1 <= len(infos) <= MAX_ITEMS or len(set(n.casefold() for n in names)) != len(names)
            or "[Content_Types].xml" not in names or "word/document.xml" not in names):
            raise TemplateError("Invalid Word package structure")
        total = 0
        for entry in infos:
            name = entry.filename
            segments = name.rstrip("/").split("/")
            if (not name or name.startswith("/") or "\\" in name or
                any(part in ("", ".", "..") for part in segments) or
                (name.endswith("/") and entry.file_size != 0) or
                ":" in name or (entry.external_attr >> 16) & 0o170000 == 0o120000):
                raise TemplateError("Unsafe Word package path")
            if entry.flag_bits & 1:
                raise TemplateError("Encrypted Word parts are not supported")
            total += entry.file_size
            if (total > MAX_UNPACKED or entry.file_size > MAX_XML and
                not name.startswith("word/media/")):
                raise TemplateError("Word package exceeds safe extraction limits")
            if (entry.file_size > 0 and
                entry.file_size > max(1024, entry.compress_size) * 250):
                raise TemplateError("Suspicious Word package compression ratio")
            low = name.casefold()
            if any(part in low for part in _BLOCKED):
                raise TemplateError("Active or embedded Word content is not accepted")
            if name.lower().endswith((".bin", ".exe", ".dll", ".js", ".vbs")):
                raise TemplateError("Executable or binary Word parts are not accepted")
            if low.endswith((".xml", ".rels")) or name == "[Content_Types].xml":
                raw = archive.read(name)
                if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
                    raise TemplateError("DTD/entities are forbidden in Word XML")
                try:
                    root = ET.fromstring(raw)
                except ET.ParseError as exc:
                    raise TemplateError("Malformed Word XML") from exc
                if low.endswith(".rels"):
                    for relationship in root:
                        reltype = relationship.attrib.get("Type", "").casefold()
                        # OPC relationship URIs normally contain /package/;
                        # inspect the terminal relationship *kind*, not the
                        # URI namespace, or every valid Word file is rejected.
                        kind = reltype.rsplit("/", 1)[-1]
                        if (relationship.attrib.get("TargetMode", "").casefold() == "external"
                            or any(part in kind for part in _DANGEROUS_REL)):
                            raise TemplateError("External or active Word relationship rejected")
                if name == "[Content_Types].xml":
                    for node in root:
                        ctype = node.attrib.get("ContentType", "").casefold()
                        if ("macroenabled" in ctype or "vba" in ctype or
                            "activex" in ctype or "oleobject" in ctype):
                            raise TemplateError("Macro/active Word content rejected")
            elif not name.startswith("word/media/") and not name.endswith("/"):
                warnings.append("Unsupported package part: " + name[:100])
        # The library may parse with or without a .dotx main content type.
        raw_types = archive.read("[Content_Types].xml")
        if b"template.main+xml" in raw_types:
            # A strict in-memory .dotx -> .docx main-type normalization is
            # required by python-docx; the original package is never changed.
            rewrites = {
                b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml":
                b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
            }
            mem = BytesIO()
            with ZipFile(mem, "w", compression=ZIP_DEFLATED) as normalized:
                for item in infos:
                    payload = archive.read(item.filename)
                    if item.filename == "[Content_Types].xml":
                        for source, target in rewrites.items():
                            payload = payload.replace(source, target)
                    normalized.writestr(item.filename, payload)
            return mem.getvalue(), tuple(sorted(set(warnings)))
    return blob, tuple(sorted(set(warnings)))


def import_word_template(path: str | Path) -> ImportResult:
    """Inspect a local Word package, then project supported static layout fields."""
    path = Path(path)
    if path.suffix.lower() not in (".docx", ".dotx"):
        raise TemplateError("Select a DOCX or DOTX file")
    if not path.is_file() or path.stat().st_size > MAX_SOURCE:
        raise TemplateError("Word source missing or exceeds the 24 MiB limit")
    try:
        blob, warnings_in = _inspect_package(path.read_bytes())
    except (BadZipFile, EOFError, RuntimeError, NotImplementedError, OSError) as exc:
        raise TemplateError("Corrupt or unsupported Word archive") from exc
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm
    try:
        doc = Document(BytesIO(blob))
    except Exception as exc:
        # python-docx/lxml may raise several provider-specific OPC/XML errors.
        # None should escape as an unexpected GUI exception or fetch a URL.
        raise TemplateError("Unsupported or corrupt Word document") from exc
    warnings = list(warnings_in)
    if len(doc.sections) != 1:
        warnings.append("Only the first section's geometry/header/footer is imported")
    sec = doc.sections[0]
    if sec.different_first_page_header_footer:
        warnings.append("Different first-page header/footer is not imported")
    if doc.settings.odd_and_even_pages_header_footer:
        warnings.append("Odd/even header and footer variants are not imported")
    if doc.paragraphs or doc.tables:
        warnings.append("Source body paragraphs/tables are not imported as report content")
    def plain(part):
        return " ".join(p.text.strip() for p in part.paragraphs if p.text.strip())
    header = plain(sec.header)
    footer = plain(sec.footer)
    if len(header) > 120 or len(footer) > 120:
        raise TemplateError("Header/footer text exceeds supported length")
    logo = b""
    for paragraph in sec.header.paragraphs:
        for run in paragraph.runs:
            for blip in run._r.xpath(".//a:blip"):
                rid = blip.get(qn("r:embed"))
                related = sec.header.part.related_parts.get(rid) if rid else None
                if related is None or not hasattr(related, "blob"):
                    raise TemplateError("Invalid embedded header image")
                raw = related.blob
                if not logo:
                    logo = raw
                else:
                    warnings.append("Additional header artwork is not imported")
    if sec.header._element.xpath(".//wp:anchor") or sec.header._element.xpath(".//w:pict"):
        warnings.append("Floating header artwork position is not imported (logo becomes inline)")
    if sec.header._element.xpath(".//w:tbl") or sec.footer._element.xpath(".//w:tbl"):
        warnings.append("Header/footer tables are not imported; rebuild them in the layout editor")
    if sec.footer._element.xpath(".//a:blip"):
        warnings.append("Footer artwork is not imported")
    if doc.element.body.xpath(".//a:blip"):
        warnings.append("Body drawings/photos are not imported as layout")
    if doc.element.body.xpath(".//w:txbxContent"):
        warnings.append("Floating text boxes are not imported")
    if doc.element.body.xpath(".//w:tbl"):
        warnings.append("Word body tables are not imported as layout")
    normal = doc.styles["Normal"]
    heading = doc.styles["Heading 1"]
    template = DocumentTemplate(
        page_width_cm=round(sec.page_width / Cm(1), 3),
        page_height_cm=round(sec.page_height / Cm(1), 3),
        left_margin_cm=round(sec.left_margin / Cm(1), 3),
        right_margin_cm=round(sec.right_margin / Cm(1), 3),
        top_margin_cm=round(sec.top_margin / Cm(1), 3),
        bottom_margin_cm=round(sec.bottom_margin / Cm(1), 3),
        header_distance_cm=round(sec.header_distance / Cm(1), 3),
        footer_distance_cm=round(sec.footer_distance / Cm(1), 3),
        header_text=header, footer_text=footer,
        normal_font=normal.font.name or "",
        normal_size_pt=round(normal.font.size.pt, 2) if normal.font.size else 11.0,
        heading_font=heading.font.name or "",
        heading_size_pt=round(heading.font.size.pt, 2) if heading.font.size else 16.0,
        logo=logo,
    )
    validate_template(template)
    if "Page" in footer:
        warnings.append("Word dynamic page-number fields are replaced with ReportMaker PAGE fields")
    if normal.font.name:
        warnings.append("PDF embeds a supported local font where available; imported Word font may be substituted")
    return ImportResult(template, tuple(dict.fromkeys(warnings)))


def save_template(template: DocumentTemplate, destination: str | Path) -> Path:
    """Atomic local sidecar. Never save original DOCX body, paths or relationships."""
    validate_template(template)
    path = Path(destination)
    if path.suffix.lower() != ".rmtemplate" or not path.parent.is_dir():
        raise TemplateError("Template requires a .rmtemplate file in an existing folder")
    record = asdict(template)
    record["logo"] = b64encode(template.logo).decode("ascii")
    payload = json.dumps({"format": FORMAT, "version": VERSION, "layout": record},
                         ensure_ascii=False, sort_keys=True).encode("utf-8")
    if len(payload) > MAX_TEMPLATE:
        raise TemplateError("Template data exceeds size limit")
    fd, tmp = tempfile.mkstemp(prefix=".reportmaker-template-", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as out:
            out.write(payload)
            out.flush()
            os.fsync(out.fileno())
        os.replace(tmp, path)
    finally:
        Path(tmp).unlink(missing_ok=True)
    return path


def load_template(source: str | Path) -> DocumentTemplate:
    path = Path(source)
    if path.suffix.lower() != ".rmtemplate" or not path.is_file():
        raise TemplateError("Template file missing")
    if path.stat().st_size > MAX_TEMPLATE:
        raise TemplateError("Template exceeds size limit")
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
        if (not isinstance(raw, dict) or set(raw) != {"format", "version", "layout"}
            or raw["format"] != FORMAT or type(raw["version"]) is not int
            or raw["version"] != VERSION or not isinstance(raw["layout"], dict)):
            raise TemplateError("Unsupported template schema")
        layout = dict(raw["layout"])
        if set(layout) != set(DocumentTemplate.__dataclass_fields__):
            raise TemplateError("Unexpected template properties")
        encoded = layout["logo"]
        if not isinstance(encoded, str) or len(encoded) > MAX_LOGO * 2:
            raise TemplateError("Invalid logo encoding")
        layout["logo"] = b64decode(encoded, validate=True)
        return validate_template(DocumentTemplate(**layout))
    except (UnicodeError, json.JSONDecodeError, TypeError, KeyError, ValueError) as exc:
        raise TemplateError("Malformed or incompatible template") from exc
